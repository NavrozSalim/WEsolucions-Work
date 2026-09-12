"""Wallkoala Excel/CSV feed: SKU + vendor price + inventory. Shipping is ignored."""
from __future__ import annotations

import csv
import io
import logging
import os
import re
import tempfile
from pathlib import Path
from typing import BinaryIO

logger = logging.getLogger("scrapers.wallkoala")

_HEADER_RE = re.compile(r"[^a-z0-9]+")
SKU_HEADERS = frozenset({"sku", "vendorsku", "vendorid", "productsku", "itemsku"})
PRICE_HEADERS = frozenset({"vendorprice", "cost", "unitprice", "buyprice", "price"})
INV_HEADERS = frozenset({
    "vendorinventory",
    "inventory",
    "qty",
    "quantity",
    "stock",
    "availableinventory",
})
SKIP_HEADERS = frozenset({"shippingprice", "shipping", "freight", "shippingcost"})
# Screenshot layout when headers are missing or unrecognized.
DEFAULT_SKU_COL = 0
DEFAULT_PRICE_COL = 1
DEFAULT_INV_COL = 3


def is_wallkoala_vendor_code(code: str | None) -> bool:
    compact = (code or "").strip().lower().replace("-", "").replace("_", "").replace(" ", "")
    return compact.startswith("wallkoala")


def _canon_header(value) -> str:
    return _HEADER_RE.sub("", str(value or "").strip().lower())


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _to_price(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("$", "").replace(",", "")
    if not text:
        return None
    try:
        return float(text)
    except (TypeError, ValueError):
        return None


def _to_stock(value) -> int:
    if value is None or value == "":
        return 0
    try:
        return max(0, int(float(value)))
    except (TypeError, ValueError):
        return 0


def _header_indexes(header_row) -> tuple[int | None, int | None, int | None]:
    sku_i = price_i = inv_i = None
    for i, cell in enumerate(header_row or []):
        key = _canon_header(cell)
        if not key or key in SKIP_HEADERS:
            continue
        if sku_i is None and key in SKU_HEADERS:
            sku_i = i
        elif price_i is None and key in PRICE_HEADERS:
            price_i = i
        elif inv_i is None and key in INV_HEADERS:
            inv_i = i
    return sku_i, price_i, inv_i


def _looks_like_header(row) -> bool:
    sku_i, price_i, inv_i = _header_indexes(row)
    return sku_i is not None and (price_i is not None or inv_i is not None)


def build_wallkoala_feed_from_rows(rows) -> dict[str, dict]:
    """SKU → {price, inventory}. Duplicate SKUs keep the last row. Shipping is skipped."""
    rows = [tuple(r) if r is not None else () for r in (rows or [])]
    if not rows:
        return {}
    start = 0
    sku_i, price_i, inv_i = _header_indexes(rows[0])
    if _looks_like_header(rows[0]):
        start = 1
    else:
        sku_i, price_i, inv_i = DEFAULT_SKU_COL, DEFAULT_PRICE_COL, DEFAULT_INV_COL

    feed: dict[str, dict] = {}
    for row in rows[start:]:
        if not row:
            continue
        sku = _cell_text(row[sku_i] if sku_i is not None and sku_i < len(row) else "")
        if not sku or sku.lower() == "sku":
            continue
        price = None
        if price_i is not None and price_i < len(row):
            price = _to_price(row[price_i])
        stock = 0
        if inv_i is not None and inv_i < len(row):
            stock = _to_stock(row[inv_i])
        feed[sku] = {"price": price, "inventory": stock, "sku": sku}
        feed[sku.lower()] = feed[sku]
    return feed


def lookup_wallkoala_entry(feed: dict | None, *keys: str) -> dict | None:
    if not feed:
        return None
    seen: set[str] = set()
    for raw in keys:
        key = str(raw or "").strip()
        if not key or key.lower() in seen:
            continue
        seen.add(key.lower())
        hit = feed.get(key) or feed.get(key.lower())
        if hit:
            return hit
    return None


def load_wallkoala_feed_from_path(path: str | Path) -> dict[str, dict]:
    path = Path(path)
    name = path.name.lower()
    if name.endswith(".csv"):
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))
        return build_wallkoala_feed_from_rows(rows)

    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    return build_wallkoala_feed_from_rows(rows)


def load_wallkoala_feed_from_file(file_obj: BinaryIO) -> dict[str, dict]:
    suffix = ".xlsx"
    name = str(getattr(file_obj, "name", "") or "").lower()
    if name.endswith(".csv"):
        suffix = ".csv"
    elif name.endswith(".xls"):
        suffix = ".xls"
    tmp = tempfile.NamedTemporaryFile(prefix="wallkoala_", suffix=suffix, delete=False)
    try:
        if hasattr(file_obj, "open"):
            try:
                file_obj.open("rb")
            except Exception:
                pass
        if hasattr(file_obj, "seek"):
            try:
                file_obj.seek(0)
            except Exception:
                pass
        for chunk in file_obj.chunks() if hasattr(file_obj, "chunks") else [file_obj.read()]:
            if chunk:
                tmp.write(chunk)
        tmp.close()
        return load_wallkoala_feed_from_path(tmp.name)
    finally:
        try:
            os.unlink(tmp.name)
        except OSError:
            pass
