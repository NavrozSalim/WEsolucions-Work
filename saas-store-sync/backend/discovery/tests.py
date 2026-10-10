import io
import os
import tempfile
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.test import TestCase, override_settings
from openpyxl import load_workbook
from rest_framework.test import APIClient

from discovery.columns import ID_FIELD, IMAGE_COLUMNS, output_columns
from discovery.engine import _live_category_rows, execute_job
from discovery.files import template_bytes, workbook_bytes
from discovery.identity import dedupe_rows
from discovery.models import DiscoveryJob
from discovery.parsers import (
    amazon_reported_total,
    amazon_results_url,
    amazon_with_price,
    listing_follow_urls,
    parse_amazon_category,
    parse_amazon_product,
    parse_ebay_category,
    parse_ebay_product,
)
from discovery.routing import QUEUE_DISCOVER_AU, QUEUE_DISCOVER_US, queue_for_marketplace
from discovery.rules import apply_rules

User = get_user_model()

STRICT_RULES = {
    'rating': {'op': 'lt', 'value': 3.5},
    'reviews': {'op': 'lt', 'value': 10},
    'price': {'op': 'gt', 'value': 80},
    'exclude_categories': ['Books'],
}


class RuleAndDedupeTests(TestCase):
    def test_rules_drop_low_rating_reviews_high_price_and_category(self):
        rows = [
            {'asin': 'B0SAMPLE01', 'rating': 4.6, 'review_count': 320, 'price': 49.99, 'category': 'Kitchen'},
            {'asin': 'B0SAMPLE02', 'rating': 3.1, 'review_count': 40, 'price': 8.5, 'category': 'Electronics'},
            {'asin': 'B0SAMPLE03', 'rating': 4.8, 'review_count': 6, 'price': 25, 'category': 'Home'},
            {'asin': 'B0SAMPLE04', 'rating': 4.4, 'review_count': 90, 'price': 189, 'category': 'Home'},
            {'asin': 'B0SAMPLE05', 'rating': 4.2, 'review_count': 200, 'price': 14, 'category': 'Books'},
        ]
        kept, removed = apply_rules(rows, STRICT_RULES)
        self.assertEqual(removed, 4)
        self.assertEqual([row['asin'] for row in kept], ['B0SAMPLE01'])

    def test_missing_rating_is_kept(self):
        kept, removed = apply_rules(
            [{'item_id': '100000000001', 'rating': None, 'price': 20, 'category': 'Kitchen'}],
            {'rating': {'op': 'lt', 'value': 3.5}},
        )
        self.assertEqual(removed, 0)
        self.assertEqual(len(kept), 1)

    def test_duplicate_asin_is_removed(self):
        rows = [
            {'asin': 'B0SAMPLE01', 'url': 'https://www.amazon.com/dp/B0SAMPLE01'},
            {'asin': 'B0SAMPLE01', 'url': 'https://www.amazon.com/dp/B0SAMPLE01?ref=dup'},
            {'url': 'https://www.amazon.com/gp/product/B0SAMPLE02'},
        ]
        unique, removed = dedupe_rows('amazon_us', rows)
        self.assertEqual(removed, 1)
        self.assertEqual(len(unique), 2)

    def test_queues_follow_region(self):
        self.assertEqual(queue_for_marketplace('amazon_us'), QUEUE_DISCOVER_US)
        self.assertEqual(queue_for_marketplace('ebay_us'), QUEUE_DISCOVER_US)
        self.assertEqual(queue_for_marketplace('amazon_au'), QUEUE_DISCOVER_AU)
        self.assertEqual(queue_for_marketplace('ebay_au'), QUEUE_DISCOVER_AU)

    def test_images_column_expands_to_ten_slots(self):
        columns = output_columns('amazon_us', ['title', 'images', 'brand'])
        self.assertEqual(columns[:4], ['asin', 'url', 'title', 'image-01'])
        self.assertEqual(columns[3:13], list(IMAGE_COLUMNS))
        self.assertEqual(columns[-1], 'brand')
        self.assertNotIn('images', columns)


class ParserTests(TestCase):
    def test_amazon_category_card(self):
        html = '''
        <div data-asin="B0SAMPLE01" data-component-type="s-search-result">
          <h2><a href="/dp/B0SAMPLE01"><span>Steel Cookware Set</span></a></h2>
          <span class="a-price"><span class="a-offscreen">$49.99</span></span>
          <span class="a-icon-alt">4.6 out of 5 stars</span>
          <span>320 ratings</span>
          <img src="https://example.com/cookware.jpg" />
        </div>
        <div data-asin=""></div>
        '''
        rows = parse_amazon_category(html, 'https://www.amazon.com/s?k=cookware')
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['asin'], 'B0SAMPLE01')
        self.assertEqual(rows[0]['url'], 'https://www.amazon.com/dp/B0SAMPLE01')
        self.assertEqual(rows[0]['price'], 49.99)
        self.assertEqual(rows[0]['rating'], 4.6)
        self.assertEqual(rows[0]['review_count'], 320)

    def test_amazon_price_keeps_cents_and_drops_sponsored_links(self):
        html = '''
        <div data-asin="B00TTQM1VW" data-component-type="s-search-result">
          <h2><a href="https://www.amazon.com/sspa/click?ie=UTF8"><span>Shure PGA52</span></a></h2>
          <span class="a-price">
            <span class="a-price-whole">429<span class="a-price-decimal">.</span></span>
            <span class="a-price-fraction">38</span>
          </span>
        </div>
        '''
        rows = parse_amazon_category(html, 'https://www.amazon.com/s?rh=n%3A8882491011')
        self.assertEqual(rows[0]['url'], 'https://www.amazon.com/dp/B00TTQM1VW')
        self.assertEqual(rows[0]['price'], 429.38)
        self.assertNotIn('sspa', rows[0]['url'])

    def test_amazon_product_page(self):
        html = '''
        <span id="productTitle">Steel Cookware Set</span>
        <span id="acrPopover" title="4.6 out of 5 stars"></span>
        <span id="acrCustomerReviewText">320 ratings</span>
        <span class="a-price"><span class="a-offscreen">$49.99</span></span>
        <a id="bylineInfo">Brand: Northline</a>
        <div id="wayfinding-breadcrumbs_feature_div"><a>Kitchen</a></div>
        <div id="feature-bullets"><li><span class="a-list-item">Oven safe</span></li></div>
        <div id="productDescription">10-piece set.</div>
        <div id="availability"><span>Only 3 left in stock - order soon.</span></div>
        '''
        row = parse_amazon_product(html, 'https://www.amazon.com/dp/B0SAMPLE01')
        self.assertEqual(row['asin'], 'B0SAMPLE01')
        self.assertEqual(row['brand'], 'Northline')
        self.assertEqual(row['category'], 'Kitchen')
        self.assertEqual(row['inventory'], 3)
        self.assertIn('Oven safe', row['bullets'])

    def test_amazon_product_reads_the_current_page(self):
        html = '''
        <span id="productTitle">
          <span class="a-truncate-full">Steel Cookware Set With Glass Lid</span>
          <span class="a-truncate-cut" aria-hidden="true">Steel Cookware…</span>
        </span>
        <span id="acrCustomerReviewText" aria-label="1,234 Reviews">1,234</span>
        <div id="availability"><span>In Stock.</span></div>
        <div id="desktop_buybox"><span>Only 7 left in stock - order soon.</span></div>
        <div id="deliveryBlockMessage">FREE delivery Tomorrow, October 11. Order within 3 hrs 12 mins</div>
        <div id="productDescription">- Dishwasher safe - 10-piece set.</div>
        <div id="feature-bullets"><li><span class="a-list-item">- Oven safe</span></li></div>
        '''
        row = parse_amazon_product(html, 'https://www.amazon.com/dp/B0SAMPLE01')
        self.assertEqual(row['title'], 'Steel Cookware Set With Glass Lid')
        self.assertNotIn('…', row['title'])
        self.assertEqual(row['review_count'], 1234)
        self.assertEqual(row['inventory'], 7)
        self.assertEqual(row['delivery_date'], 'Tomorrow, October 11')
        self.assertEqual(row['description'], 'Dishwasher safe 10-piece set.')
        self.assertEqual(row['bullets'], 'Oven safe')
        self.assertNotIn(' - ', row['description'])

    def test_amazon_product_keeps_ten_high_res_images(self):
        hires = [
            f'https://m.media-amazon.com/images/I/71IMAGE{index:04d}.jpg'
            for index in range(1, 13)
        ]
        script = ','.join(
            '{"hiRes":"%s","thumb":"%s","large":"%s"}' % (
                url,
                url.replace('.jpg', '._AC_US40_.jpg'),
                url.replace('.jpg', '._AC_SL1500_.jpg'),
            )
            for url in hires
        )
        thumbs = ''.join(
            f'<img src="{url.replace(".jpg", "._AC_US40_.jpg")}" />'
            for url in hires
        )
        html = f'''
        <script>"colorImages":{{"initial":[{script}]}}</script>
        <div id="altImages">{thumbs}</div>
        <img id="landingImage" src="{hires[0].replace(".jpg", "._AC_SX679_.jpg")}" />
        '''
        row = parse_amazon_product(html, 'https://www.amazon.com/dp/B0SAMPLE01')
        self.assertEqual(row['image-01'], hires[0].replace('.jpg', '._US_1500_.jpg'))
        self.assertEqual(row['image-10'], hires[9].replace('.jpg', '._US_1500_.jpg'))
        self.assertNotIn(hires[10].replace('.jpg', '._US_1500_.jpg'), row.values())
        self.assertNotIn('._AC_US40_', ' '.join(row[name] for name in IMAGE_COLUMNS))
        self.assertTrue(all(not row[name] or row[name].endswith('._US_1500_.jpg') for name in IMAGE_COLUMNS))

    def test_amazon_product_upgrades_thumbnails_and_leaves_empty_slots(self):
        html = '''
        <div id="altImages">
          <img src="https://m.media-amazon.com/images/I/71ONLYTHUMB._AC_US40_.jpg" />
        </div>
        <img id="landingImage" src="https://m.media-amazon.com/images/I/71ONLYTHUMB._AC_SX38_.jpg" />
        '''
        row = parse_amazon_product(html, 'https://www.amazon.com/dp/B0SAMPLE01')
        self.assertEqual(row['image-01'], 'https://m.media-amazon.com/images/I/71ONLYTHUMB._US_1500_.jpg')
        self.assertEqual(row['image-02'], '')
        self.assertEqual(row['image-10'], '')

    def test_ebay_product_keeps_zoom_images_only(self):
        html = '''
        <div class="ux-image-carousel">
          <img src="https://i.ebayimg.com/images/g/abc/s-l64.jpg"
               data-zoom-src="https://i.ebayimg.com/images/g/abc/s-l1600.jpg" />
          <img src="https://i.ebayimg.com/images/g/def/s-l96.jpg" />
          <img src="https://i.ebayimg.com/images/g/abc/s-l500.jpg" />
        </div>
        '''
        row = parse_ebay_product(html, 'https://www.ebay.com/itm/100000000001')
        self.assertEqual(row['image-01'], 'https://i.ebayimg.com/images/g/abc/s-l1600.jpg')
        self.assertEqual(row['image-02'], 'https://i.ebayimg.com/images/g/def/s-l1600.jpg')
        self.assertEqual(row['image-03'], '')

    def test_amazon_total_and_price_band(self):
        html = '<span>1-48 of over 50,000 results</span>'
        self.assertEqual(amazon_reported_total(html), 50000)
        self.assertEqual(amazon_reported_total('<span>10,000+ results</span>'), 10000)
        self.assertEqual(
            amazon_reported_total('{"totalResultCount":41000,"other":1}'),
            41000,
        )
        self.assertEqual(
            amazon_reported_total('of over <span>60,000</span> results'),
            60000,
        )
        self.assertIsNone(amazon_reported_total('<span>No results</span>'))
        sliced = amazon_with_price(
            'https://www.amazon.com/s?rh=n%3A8882491011&fs=true&page=3',
            0,
            2500,
        )
        self.assertIn('p_36%3A0-2500', sliced)
        self.assertNotIn('page=3', sliced)
        self.assertIn('8882491011', sliced)

    def test_amazon_node_page_opens_the_full_grid(self):
        url = (
            'https://www.amazon.com/b/?ie=UTF8&node=8882491011'
            '&pf_rd_p=8d993004-e7d9-458b-9b36-18d9c6d990c8'
        )
        self.assertEqual(
            amazon_results_url(url),
            'https://www.amazon.com/s?rh=n%3A8882491011&fs=true',
        )
        html = '''
        <a href="/s?rh=n%3A8882491011&fs=true">View products</a>
        <a href="/dp/B0SAMPLE09">Featured Mixer</a>
        '''
        follows = listing_follow_urls(html, 'https://www.amazon.com/b/?node=8882491011')
        self.assertEqual(follows, ['https://www.amazon.com/s?rh=n%3A8882491011&fs=true'])
        rows = parse_amazon_category(html, 'https://www.amazon.com/b/?node=8882491011')
        self.assertEqual(rows[0]['asin'], 'B0SAMPLE09')
        self.assertEqual(rows[0]['title'], 'Featured Mixer')
        embedded = '<html><script>{"asin":"B0SAMPLE07"}</script><a href="/dp/B0SAMPLE08">Listed Mixer</a></html>'
        extra = parse_amazon_category(embedded, 'https://www.amazon.com/b/?node=8882491011')
        self.assertEqual([row['asin'] for row in extra], ['B0SAMPLE08'])
        self.assertEqual(extra[0]['url'], 'https://www.amazon.com/dp/B0SAMPLE08')

    def test_ebay_category_card(self):
        html = '''
        <li class="s-item">
          <a class="s-item__link" href="https://www.ebay.com/itm/100000000001?hash=abc"></a>
          <div class="s-item__title"><span>Steel Cookware Set</span></div>
          <span class="s-item__price">$49.99</span>
        </li>
        <li class="s-item">
          <a class="s-item__link" href="https://www.ebay.com/itm/0"></a>
          <div class="s-item__title">Shop on eBay</div>
        </li>
        '''
        rows = parse_ebay_category(html, 'https://www.ebay.com/sch/i.html?_nkw=cookware')
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['item_id'], '100000000001')
        self.assertEqual(rows[0]['price'], 49.99)


@override_settings(MEDIA_ROOT=tempfile.mkdtemp(prefix='discovery-test-'))
class AmazonCategoryWalkTests(TestCase):
    def test_category_walks_past_300_products(self):
        def fake_fetch(url, session=None):
            query = __import__('urllib.parse', fromlist=['parse_qs']).parse_qs(
                __import__('urllib.parse', fromlist=['urlsplit']).urlsplit(url).query,
            )
            page = int((query.get('page') or ['1'])[0])
            if page > 25:
                return '<html></html>'
            cards = []
            for offset in range(16):
                asin = f'B{(page * 16 + offset):09d}'
                cards.append(
                    f'<div data-asin="{asin}"><h2><a href="/dp/{asin}">Item {asin}</a></h2></div>'
                )
            return ''.join(cards)

        with patch('discovery.engine._fetch', side_effect=fake_fetch):
            rows = _live_category_rows(
                'https://www.amazon.com/s?rh=n%3A8882491011&fs=true',
                'amazon_us',
                [],
            )
        self.assertGreater(len(rows), 300)
        self.assertEqual(len(rows), 25 * 16)

    def test_large_amazon_category_is_split_by_price(self):
        def fake_fetch(url, session=None):
            if 'p_36' in url:
                return '<html></html>'
            query = __import__('urllib.parse', fromlist=['parse_qs']).parse_qs(
                __import__('urllib.parse', fromlist=['urlsplit']).urlsplit(url).query,
            )
            page = int((query.get('page') or ['1'])[0])
            if page > 1:
                return '<html></html>'
            return (
                '<span>1-1 of over 10,000 results</span>'
                '<div data-asin="B0SAMPLE01"><h2><a href="/dp/B0SAMPLE01">Mixer</a></h2></div>'
            )

        with patch('discovery.engine._fetch', side_effect=fake_fetch) as fetch:
            rows = _live_category_rows(
                'https://www.amazon.com/s?rh=n%3A8882491011&fs=true',
                'amazon_us',
                [],
            )
        self.assertEqual([row['asin'] for row in rows], ['B0SAMPLE01'])
        self.assertGreaterEqual(sum(1 for call in fetch.call_args_list if 'p_36' in call.args[0]), 2)

    def test_repeated_amazon_page_is_split_by_price(self):
        """Amazon repeats the same cards after a few hundred. That is a cap, not the end."""
        first = ''.join(
            f'<div data-asin="B0SAMPLE{n:02d}"><h2><a href="/dp/B0SAMPLE{n:02d}">Item {n}</a></h2></div>'
            for n in range(1, 3)
        )

        def fake_fetch(url, session=None):
            if 'p_36' in url:
                if 'page=2' in url or 'page%3D2' in url:
                    return '<html></html>'
                if 'p_36%3A0-' in url or 'p_36:0-' in url:
                    return (
                        '<div data-asin="B0LOW00001">'
                        '<h2><a href="/dp/B0LOW00001">Low</a></h2></div>'
                    )
                return (
                    '<div data-asin="B0HIGH0001">'
                    '<h2><a href="/dp/B0HIGH0001">High</a></h2></div>'
                )
            return first

        with patch('discovery.engine._fetch', side_effect=fake_fetch):
            rows = _live_category_rows(
                'https://www.amazon.com/s?rh=n%3A8882491011&fs=true',
                'amazon_us',
                [],
            )
        self.assertEqual(
            [row['asin'] for row in rows],
            ['B0SAMPLE01', 'B0SAMPLE02', 'B0LOW00001', 'B0HIGH0001'],
        )

    def test_finished_amazon_category_is_not_split(self):
        def fake_fetch(url, session=None):
            if 'page=2' in url or 'p_36' in url:
                return '<html></html>'
            return (
                '<div data-asin="B0SAMPLE01">'
                '<h2><a href="/dp/B0SAMPLE01">Mixer</a></h2></div>'
            )

        with patch('discovery.engine._fetch', side_effect=fake_fetch) as fetch:
            rows = _live_category_rows(
                'https://www.amazon.com/s?rh=n%3A8882491011&fs=true',
                'amazon_us',
                [],
            )
        self.assertEqual([row['asin'] for row in rows], ['B0SAMPLE01'])
        self.assertFalse(any('p_36' in call.args[0] for call in fetch.call_args_list))


class DiscoveryJobTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username='discoverer',
            email='discoverer@example.com',
            password='pass12345',
        )
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self.async_patch = patch('discovery.views.run_discovery_job.apply_async')
        self.async_patch.start()

    def tearDown(self):
        self.async_patch.stop()

    def _run_template(self, marketplace, mode, columns):
        filename, payload = template_bytes(marketplace, mode)
        response = self.client.post(
            '/api/v1/discovery/jobs/',
            {
                'marketplace': marketplace,
                'mode': mode,
                'use_sample': 'true',
                'rules': __import__('json').dumps(STRICT_RULES),
                'columns': __import__('json').dumps(columns),
                'file': ContentFile(payload, name=filename),
            },
            format='multipart',
        )
        self.assertEqual(response.status_code, 201, response.content)
        job = DiscoveryJob.objects.get(id=response.data['id'])
        self.assertEqual(job.queue_name, queue_for_marketplace(marketplace))
        execute_job(job.id)
        job.refresh_from_db()
        return job

    def test_upload_must_match_the_template(self):
        payload = workbook_bytes(['title', 'price'], [{'title': 'Mic', 'price': 10}])
        response = self.client.post(
            '/api/v1/discovery/jobs/',
            {
                'marketplace': 'amazon_us',
                'mode': 'category',
                'use_sample': 'true',
                'rules': '{}',
                'columns': '[]',
                'file': ContentFile(payload, name='wrong.xlsx'),
            },
            format='multipart',
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn('url', response.data['detail'])
        product = workbook_bytes(['url'], [{'url': 'https://www.amazon.com/s?k=mics'}])
        response = self.client.post(
            '/api/v1/discovery/jobs/',
            {
                'marketplace': 'amazon_us',
                'mode': 'product',
                'use_sample': 'true',
                'rules': '{}',
                'columns': '[]',
                'file': ContentFile(product, name='category-as-product.xlsx'),
            },
            format='multipart',
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn('asin, url', response.data['detail'])

    def test_category_sample_dedupes_and_filters(self):
        job = self._run_template(
            'amazon_us',
            'category',
            ['asin', 'title', 'price', 'rating', 'review_count', 'category'],
        )
        self.assertEqual(job.status, DiscoveryJob.Status.SUCCEEDED)
        stats = job.stats
        # Two category URLs each return 6 cards (5 products + 1 duplicate).
        self.assertEqual(stats['scraped'], 12)
        self.assertEqual(stats['duplicates_removed'], 7)
        self.assertEqual(stats['removed_by_rules'], 4)
        self.assertEqual(stats['kept'], 1)
        workbook = load_workbook(io.BytesIO(job.result_file.read()))
        sheet = workbook.active
        headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
        self.assertEqual(headers[0], 'asin')
        self.assertIn('url', headers)
        body = list(sheet.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(body), 1)
        self.assertEqual(body[0][0], 'B0SAMPLE01')

    def test_product_sample_skips_duplicate_before_scrape(self):
        job = self._run_template('ebay_au', 'product', ['title', 'price'])
        self.assertEqual(job.status, DiscoveryJob.Status.SUCCEEDED, job.error_message)
        self.assertEqual(job.queue_name, QUEUE_DISCOVER_AU)
        self.assertGreaterEqual(job.stats['duplicates_removed'], 1)
        self.assertEqual(job.stats['kept'], 1)
        workbook = load_workbook(io.BytesIO(job.result_file.read()))
        headers = [cell.value for cell in next(workbook.active.iter_rows(max_row=1))]
        self.assertEqual(headers[0], ID_FIELD['ebay_au'])
        self.assertIn('title', headers)
        self.assertNotIn('description', headers)

    def test_selected_images_fill_ten_columns(self):
        filename, payload = template_bytes('amazon_us', 'product')
        response = self.client.post(
            '/api/v1/discovery/jobs/',
            {
                'marketplace': 'amazon_us',
                'mode': 'product',
                'use_sample': 'true',
                'rules': '{}',
                'columns': __import__('json').dumps(['images']),
                'file': ContentFile(payload, name=filename),
            },
            format='multipart',
        )
        self.assertEqual(response.status_code, 201, response.content)
        job = DiscoveryJob.objects.get(id=response.data['id'])
        execute_job(job.id)
        job.refresh_from_db()
        self.assertEqual(job.status, DiscoveryJob.Status.SUCCEEDED, job.error_message)
        workbook = load_workbook(io.BytesIO(job.result_file.read()))
        headers = [cell.value for cell in next(workbook.active.iter_rows(max_row=1))]
        self.assertEqual(headers[2:12], list(IMAGE_COLUMNS))
        self.assertNotIn('images', headers)
        first = next(workbook.active.iter_rows(min_row=2, values_only=True))
        values = dict(zip(headers, first))
        self.assertEqual(values['image-01'], 'https://example.com/cookware.jpg')
        self.assertEqual(values['image-02'], 'https://example.com/cookware-2.jpg')
        self.assertFalse(values['image-03'])
        self.assertFalse(values['image-10'])

    def test_rows_table_and_delete(self):
        job = self._run_template(
            'amazon_us',
            'category',
            ['asin', 'title', 'price'],
        )
        listed = self.client.get('/api/v1/discovery/jobs/')
        self.assertEqual(listed.status_code, 200)
        self.assertTrue(any(item['id'] == str(job.id) for item in listed.data))
        rows = self.client.get(f'/api/v1/discovery/jobs/{job.id}/rows/')
        self.assertEqual(rows.status_code, 200)
        self.assertIn('title', rows.data['columns'])
        self.assertEqual(rows.data['total'], 1)
        self.assertEqual(rows.data['rows'][0]['asin'], 'B0SAMPLE01')
        ids = self.client.get(f'/api/v1/discovery/jobs/{job.id}/ids/')
        self.assertEqual(ids.status_code, 200)
        id_sheet = load_workbook(io.BytesIO(b''.join(ids.streaming_content))).active
        id_headers = [cell.value for cell in next(id_sheet.iter_rows(max_row=1))]
        self.assertEqual(id_headers, ['asin'])
        self.assertEqual(next(id_sheet.iter_rows(min_row=2, values_only=True))[0], 'B0SAMPLE01')
        deleted = self.client.delete(f'/api/v1/discovery/jobs/{job.id}/')
        self.assertEqual(deleted.status_code, 204)
        self.assertFalse(DiscoveryJob.objects.filter(id=job.id).exists())

    def test_continue_category_into_product_details(self):
        parent = self._run_template('amazon_au', 'category', ['asin', 'url', 'title', 'price', 'rating', 'review_count', 'category'])
        response = self.client.post(
            f'/api/v1/discovery/jobs/{parent.id}/continue/',
            {'use_sample': True, 'rules': STRICT_RULES, 'columns': ['title', 'brand', 'price']},
            format='json',
        )
        self.assertEqual(response.status_code, 201, response.content)
        child = DiscoveryJob.objects.get(id=response.data['id'])
        self.assertEqual(child.mode, DiscoveryJob.Mode.PRODUCT)
        self.assertEqual(child.marketplace, 'amazon_au')
        execute_job(child.id)
        child.refresh_from_db()
        self.assertEqual(child.status, DiscoveryJob.Status.SUCCEEDED, child.error_message)
        self.assertEqual(child.stats['kept'], 1)
        self.assertEqual(child.stats['input_rows'], 1)
        self.assertTrue(child.source_bytes)

    def test_product_file_scrapes_every_unique_product(self):
        count = 201
        rows = [
            {'url': f'https://www.amazon.com/dp/B{index:09d}'}
            for index in range(1, count + 1)
        ]
        payload = workbook_bytes(['url'], rows)
        job = DiscoveryJob(
            owner=self.user,
            marketplace='amazon_us',
            mode=DiscoveryJob.Mode.PRODUCT,
            original_filename='amazon_us-from-category.xlsx',
            use_sample=True,
            zip_code='10001',
            source_bytes=payload,
        )
        job.source_file.save('products.xlsx', ContentFile(payload), save=False)
        job.save()
        execute_job(job.id)
        job.refresh_from_db()
        self.assertEqual(job.status, DiscoveryJob.Status.SUCCEEDED, job.error_message)
        self.assertEqual(job.error_message, '')
        self.assertEqual(job.stats['kept'], count)
        self.assertEqual(job.stats['warning'], '')
        self.assertEqual(job.products.count(), count)

    def test_scrape_reads_the_database_copy_when_the_disk_file_is_missing(self):
        filename, payload = template_bytes('amazon_us', 'product')
        response = self.client.post(
            '/api/v1/discovery/jobs/',
            {
                'marketplace': 'amazon_us',
                'mode': 'product',
                'use_sample': 'true',
                'rules': '{}',
                'columns': '[]',
                'file': ContentFile(payload, name=filename),
            },
            format='multipart',
        )
        self.assertEqual(response.status_code, 201, response.content)
        job = DiscoveryJob.objects.get(id=response.data['id'])
        self.assertTrue(job.source_bytes)
        os.remove(job.source_file.path)
        execute_job(job.id)
        job.refresh_from_db()
        self.assertEqual(job.status, DiscoveryJob.Status.SUCCEEDED, job.error_message)
        self.assertTrue(job.result_bytes)
        os.remove(job.result_file.path)
        download = self.client.get(f'/api/v1/discovery/jobs/{job.id}/download/')
        self.assertEqual(download.status_code, 200)
        sheet = load_workbook(io.BytesIO(b''.join(download.streaming_content))).active
        self.assertEqual([cell.value for cell in next(sheet.iter_rows(max_row=1))][0], 'asin')
