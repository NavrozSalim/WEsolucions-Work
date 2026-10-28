"""Upload and result spreadsheets."""

from __future__ import annotations

import csv
import io
import re
from urllib.parse import urlsplit

from openpyxl import Workbook, load_workbook

from .columns import HOSTS, ID_FIELD, is_amazon
from .identity import extract_asin, extract_ebay_item_id

# Excel rejects these control characters and any cell longer than this.
_ILLEGAL_XLSX = re.compile(r'[\000-\010]|[\013-\014]|[\016-\037]')
_XLSX_CELL_LIMIT = 32767


def _sheet_value(value):
    """Value openpyxl can write. Scraped text often includes characters Excel rejects."""
    if value is None or value == '':
        return ''
    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if value != value or value in (float('inf'), float('-inf')):
            return ''
        return value
    if not isinstance(value, str):
        value = str(value)
    value = _ILLEGAL_XLSX.sub('', value)
    if len(value) > _XLSX_CELL_LIMIT:
        value = value[:_XLSX_CELL_LIMIT]
    return value


_HEADER_ALIASES = {
    'url': 'url',
    'link': 'url',
    'product_url': 'url',
    'product url': 'url',
    'category_url': 'url',
    'category url': 'url',
    'vendor_url': 'url',
    'vendor url': 'url',
    'asin': 'asin',
    'item_id': 'item_id',
    'item id': 'item_id',
    'itemid': 'item_id',
    'ebay_item_id': 'item_id',
    'title': 'title',
    'price': 'price',
    'rating': 'rating',
    'stars': 'rating',
    'review_count': 'review_count',
    'reviews': 'review_count',
    'review count': 'review_count',
    'category': 'category',
    'brand': 'brand',
}


class SpreadsheetError(Exception):
    pass


def _header_key(value) -> str:
    text = re.sub(r'\s+', ' ', str(value or '').strip().lower())
    return _HEADER_ALIASES.get(text, '')


def template_columns(marketplace: str, mode: str) -> list[str]:
    """Header row of the downloadable template for this marketplace and mode."""
    if mode == 'category':
        return ['url']
    return [ID_FIELD[marketplace], 'url']


def _rows_from_matrix(matrix: list[list], expected: list[str] | None = None) -> list[dict]:
    if not matrix:
        raise SpreadsheetError('The file is empty.')
    if expected is not None:
        actual = [str(cell or '').strip().lower() for cell in matrix[0]]
        while actual and actual[-1] == '':
            actual.pop()
        if actual != [column.lower() for column in expected]:
            shown = ', '.join(expected)
            raise SpreadsheetError(
                f'This file does not match the template. The first row must be: {shown}.'
            )
        headers = list(expected)
    else:
        headers = [_header_key(cell) for cell in matrix[0]]
    if 'url' not in headers and 'asin' not in headers and 'item_id' not in headers:
        raise SpreadsheetError(
            'The file needs a url column, or an asin / item_id column.'
        )
    rows = []
    for raw in matrix[1:]:
        if not raw or not any(str(cell or '').strip() for cell in raw):
            continue
        row = {}
        for index, key in enumerate(headers):
            if not key or index >= len(raw):
                continue
            value = raw[index]
            if value is None:
                continue
            row[key] = value if isinstance(value, (int, float)) else str(value).strip()
        if row.get('url') or row.get('asin') or row.get('item_id'):
            rows.append(row)
    if not rows:
        raise SpreadsheetError('The file has headers but no data rows.')
    return rows


def open_bytes(data: bytes, name: str):
    """File-like object so a database copy still has a filename."""
    handle = io.BytesIO(data)
    handle.name = name or 'upload.xlsx'
    return handle


def read_spreadsheet(uploaded, marketplace: str | None = None, mode: str | None = None) -> list[dict]:
    name = (getattr(uploaded, 'name', '') or '').lower()
    data = uploaded.read()
    if hasattr(uploaded, 'seek'):
        uploaded.seek(0)
    expected = template_columns(marketplace, mode) if marketplace and mode else None
    if name.endswith('.csv'):
        text = data.decode('utf-8-sig', errors='replace')
        reader = csv.reader(io.StringIO(text))
        return _rows_from_matrix([row for row in reader], expected)
    if name.endswith('.xlsx') or name.endswith('.xls'):
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.active
        matrix = [list(row) for row in sheet.iter_rows(values_only=True)]
        workbook.close()
        return _rows_from_matrix(matrix, expected)
    raise SpreadsheetError('Upload a CSV or XLSX file.')


def product_url(marketplace: str, row: dict) -> str:
    url = str(row.get('url') or '').strip()
    host = HOSTS[marketplace]
    if not is_amazon(marketplace):
        item_id = extract_ebay_item_id(str(row.get('item_id') or '')) or extract_ebay_item_id(url)
        if item_id:
            return f'{host}/itm/{item_id}'
    if url.startswith('http://') or url.startswith('https://'):
        return url
    if is_amazon(marketplace):
        asin = extract_asin(str(row.get('asin') or '')) or extract_asin(url)
        if asin:
            return f'{host}/dp/{asin}'
    return url


def stamp_identity(marketplace: str, row: dict) -> dict:
    out = dict(row)
    url = product_url(marketplace, out)
    if url:
        out['url'] = url
    if is_amazon(marketplace):
        asin = extract_asin(str(out.get('asin') or '')) or extract_asin(url)
        if asin:
            out['asin'] = asin
    else:
        item_id = extract_ebay_item_id(str(out.get('item_id') or '')) or extract_ebay_item_id(url)
        if item_id:
            out['item_id'] = item_id
    return out


def looks_like_category_url(url: str) -> bool:
    path = urlsplit(url).path.lower()
    if '/dp/' in path or '/itm/' in path or '/gp/product/' in path:
        return False
    return True


def read_result_sheet(uploaded) -> tuple[list[str], list[dict]]:
    """Read a result workbook, keeping every column the job wrote."""
    data = uploaded.read() if hasattr(uploaded, 'read') else uploaded
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    matrix = [list(row) for row in sheet.iter_rows(values_only=True)]
    workbook.close()
    if not matrix:
        return [], []
    columns = [str(cell).strip() for cell in matrix[0] if str(cell or '').strip()]
    rows = []
    for raw in matrix[1:]:
        if not raw or not any(cell not in (None, '') for cell in raw):
            continue
        row = {}
        for index, column in enumerate(columns):
            value = raw[index] if index < len(raw) else ''
            row[column] = '' if value is None else value
        rows.append(row)
    return columns, rows


def csv_chunks(columns: list[str], rows, batch: int = 200):
    """CSV bytes in small pieces, including the Excel BOM on the first piece.

    Callers can stream a large result without holding the whole file in memory.
    """
    columns = list(columns)
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=columns, extrasaction='ignore')
    writer.writeheader()
    pending = 0
    started = False

    def flush(with_bom: bool):
        text = buffer.getvalue()
        buffer.seek(0)
        buffer.truncate(0)
        if not text:
            return b''
        payload = text.encode('utf-8')
        if with_bom:
            return b'\xef\xbb\xbf' + payload
        return payload

    for row in rows:
        data = row or {}
        writer.writerow({
            column: _sheet_value(data.get(column, ''))
            for column in columns
        })
        pending += 1
        if pending >= batch:
            payload = flush(not started)
            started = True
            pending = 0
            if payload:
                yield payload
    payload = flush(not started)
    if payload:
        yield payload


def csv_bytes(columns: list[str], rows) -> bytes:
    return b''.join(csv_chunks(columns, rows))


def write_workbook(columns: list[str], rows, handle) -> None:
    """Write an xlsx straight to ``handle``.

    ``write_only`` flushes each row instead of keeping every cell in memory,
    which is what made an 80,000-row Amazon sheet blow the web worker.
    """
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet('Results')
    columns = list(columns)
    sheet.append([_sheet_value(column) for column in columns])
    for row in rows:
        data = row or {}
        sheet.append([_sheet_value(data.get(column, '')) for column in columns])
    workbook.save(handle)


def workbook_bytes(columns: list[str], rows) -> bytes:
    buffer = io.BytesIO()
    write_workbook(columns, rows, buffer)
    return buffer.getvalue()


def template_bytes(marketplace: str, mode: str) -> tuple[str, bytes]:
    from .sample_data import category_template_urls, product_template_rows

    field = ID_FIELD[marketplace]
    if mode == 'category':
        columns = ['url']
        rows = [{'url': url} for url in category_template_urls(marketplace)]
        filename = f'{marketplace}-category-template.xlsx'
    else:
        columns = [field, 'url']
        rows = [
            {field: row[field], 'url': row['url']}
            for row in product_template_rows(marketplace)
        ]
        filename = f'{marketplace}-product-template.xlsx'
    return filename, workbook_bytes(columns, rows)
