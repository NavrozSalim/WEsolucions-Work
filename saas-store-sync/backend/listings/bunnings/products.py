"""Map managed StoreListing rows to Mirakl product (P41) + offer (OF01) CSVs."""
from __future__ import annotations

import csv
import io
import json
import logging
import re
import threading
import time
from decimal import Decimal, InvalidOperation

from django.utils import timezone

from ..errors import MarketplaceError
from ..models import ListingAction, ListingStatus, StoreListing
from .client import BunningsClient, extract_import_id, line_error_rows

logger = logging.getLogger("listings.bunnings")

PRODUCT_WAIT_ATTEMPTS = 6
PRODUCT_WAIT_SECONDS = 2.0
OFFER_RETRY_ATTEMPTS = 4
OFFER_RETRY_SECONDS = 3.0
_PRODUCT_MISSING_MARKERS = (
    "product does not exist",
    "the product does not exist",
    "unknown product",
    "product not found",
    "product is not found",
)
PENDING_OFFER_MESSAGE = (
    "Product was sent to Bunnings and is waiting for review. "
    "Publish again after Bunnings accepts it to create the offer."
)

OFFER_STATE_NEW = "11"
PRODUCT_ID_TYPE_SHOP_SKU = "SHOP_SKU"
PRODUCT_CSV_HEADERS = [
    "category",
    "product-id",
    "product-id-type",
    "variant-group-code",
    "title",
    "description",
    "brand",
    "ean",
    "DISPLAY_NAME",
    "PRODUCT_DESCRIPTION",
    "LONG_DESCRIPTION",
    "BRAND",
    "GTIN",
    "PRIMARY_IMAGE",
    "CATEGORY",
]
VARIANT_GROUP_HEADER = "variant-group-code"
OPERATOR_VARIANT_GROUP = "VARIANT_GROUP_CODE"
OPTION_ATTR_ALIASES = {
    "size": "size",
    "colour": "colour",
    "color": "color",
    "finish": "finish",
    "width": "width",
    "length": "product-length",
    "height": "product-height",
}
OFFER_CSV_HEADERS = [
    "sku",
    "product-id",
    "product-id-type",
    "description",
    "price",
    "quantity",
    "state",
    "logistic-class",
    "leadtime-to-ship",
    "update-delete",
]


def _dec(value, default="0") -> Decimal:
    try:
        return Decimal(str(value if value is not None and value != "" else default))
    except (InvalidOperation, TypeError, ValueError):
        return Decimal(default)


def _photo_urls(listing_or_raw) -> list[str]:
    raw = listing_or_raw
    if hasattr(listing_or_raw, "image_urls"):
        raw = listing_or_raw.image_urls
    text = str(raw or "").strip()
    if not text:
        return []
    for sep in ("|", ";", "\n", ","):
        if sep in text:
            return [p.strip() for p in text.replace(";", "|").replace("\n", "|").replace(",", "|").split("|") if p.strip()][:9]
    return [text] if text.startswith("http") else []


# Bunnings reads these operator codes. PM11 replaces the list when a category
# defines its own additional-image attributes. image-2 is not one of them.
FALLBACK_EXTRA_IMAGE_CODES = tuple(
    (f"ADDITIONAL_IMAGE_{index}", f"Additional image {index}")
    for index in range(1, 9)
)
_EXTRA_IMAGE_LABEL_RE = re.compile(
    r"(?:additional|alt|alternative|secondary) image 0*(\d+)"
)
_PLAIN_IMAGE_RE = re.compile(r"image 0*(\d+)")
_NOT_GALLERY_IMAGE = ("brochure", "document", "swatch", "thumbnail", "icon", "logo")


def additional_image_slot(text: str) -> int | None:
    """1-based additional-image position. image-2 and Additional image 1 are both slot 1."""
    raw = str(text or "").strip().lower().replace("-", " ").replace("_", " ")
    raw = re.sub(r"\s+", " ", raw).strip()
    raw = re.sub(r"\(optional\)$", "", raw).strip()
    if not raw or "primary" in raw:
        return None
    if any(word in raw for word in _NOT_GALLERY_IMAGE):
        return None
    labelled = _EXTRA_IMAGE_LABEL_RE.fullmatch(raw)
    if labelled:
        number = int(labelled.group(1))
        return number if number >= 1 else None
    plain = _PLAIN_IMAGE_RE.fullmatch(raw)
    if plain:
        number = int(plain.group(1))
        if number <= 1:
            return None
        return number - 1
    return None


def is_additional_image_attribute(code: str, label: str = "", attr_type: str = "") -> bool:
    if not (additional_image_slot(code) or additional_image_slot(label)):
        return False
    kind = str(attr_type or "").strip().upper()
    if kind in {"LIST", "INTEGER", "DECIMAL", "NUMERIC", "BOOLEAN", "DATE", "MEASUREMENT"}:
        return False
    return True


def extra_image_codes_from_items(items) -> list[str]:
    """Operator codes for additional images, ordered by image slot."""
    found: list[tuple[int, str]] = []
    seen: set[str] = set()
    for item in items or []:
        if not isinstance(item, dict):
            continue
        code = str(item.get("code") or "").strip()
        if not code or code.lower() in seen:
            continue
        label = str(item.get("label") or "")
        kind = str(item.get("type") or "")
        if not item.get("extra_image") and not is_additional_image_attribute(code, label, kind):
            continue
        slot = additional_image_slot(code) or additional_image_slot(label) or (len(found) + 1)
        seen.add(code.lower())
        found.append((slot, code))
    found.sort(key=lambda pair: (pair[0], pair[1].lower()))
    return [code for _slot, code in found]


def extra_image_codes_for(store, categories) -> list[str]:
    items = []
    if store is not None:
        for raw in categories or []:
            category = str(raw or "").strip()
            if category:
                items.extend(load_category_attributes(store, category))
    codes = extra_image_codes_from_items(items)
    if codes:
        return codes
    return [code for code, _label in FALLBACK_EXTRA_IMAGE_CODES]


def gallery_urls(listing, attrs: dict | None = None) -> list[str]:
    """Primary image first, then additional images from the listing and saved attributes."""
    photos = _photo_urls(listing)
    variant_img = str(getattr(listing, "variation_image_url", "") or "").strip()
    if variant_img:
        photos = [variant_img] + [url for url in photos if url != variant_img]
    ordered = list(photos)
    slots: list[tuple[int, str]] = []
    for code, value in (attrs or {}).items():
        slot = additional_image_slot(code)
        url = str(value or "").strip()
        if slot and url.startswith("http"):
            slots.append((slot, url))
    for _slot, url in sorted(slots, key=lambda pair: pair[0]):
        if url not in ordered:
            ordered.append(url)
    return ordered[:9]


def parse_extras(listing_or_json) -> dict:
    raw = listing_or_json
    if hasattr(listing_or_json, "external_data_object_json"):
        raw = listing_or_json.external_data_object_json
    if isinstance(raw, dict):
        data = raw
    elif isinstance(raw, str) and raw.strip():
        try:
            data = json.loads(raw)
        except (TypeError, ValueError, json.JSONDecodeError):
            data = {}
    else:
        data = {}
    if not isinstance(data, dict):
        return {}
    if data.get("marketplace") and str(data.get("marketplace")).lower() != "bunnings":
        return {}
    return data


CORE_PRODUCT_CODES = frozenset({
    "category",
    "product-id",
    "product-id-type",
    "variant-group-code",
    "title",
    "description",
    "brand",
    "ean",
    "sku",
    "image-1",
})
_ATTR_CACHE: dict[tuple, list] = {}
_ATTR_LOCK = threading.Lock()
_VALUE_LIST_CACHE: dict[str, dict[str, list[dict]]] = {}
# Mirakl LIST attributes must be sent as value-list codes, not labels.
DEFAULT_ATTR_VALUE_LISTS = {
    "BRAND": "BRAND_LIST",
    "PRIMARY_UOM": "UNITS_OF_MEASURE",
}
MAX_INLINE_LIST_VALUES = 400
_DIM_FALLBACKS = {
    "weight": ("weight", "product-weight", "gross-weight"),
    "length": ("length", "product-length", "package-length"),
    "height": ("height", "product-height", "package-height"),
    "width": ("width", "product-width", "package-width"),
}
# Already present as Mirakl identification / offer columns on the bulk template.
TEMPLATE_SKIP_ATTRS = frozenset({
    "category",
    "product-id",
    "product-id-type",
    "variant-group-code",
    "variant_group_code",
    "sku",
    "price",
    "quantity",
    "state",
    "logistic-class",
    "leadtime-to-ship",
    "update-delete",
    "update_delete",
    "ean",
    "title",
    "description",
    "image-1",
    "image-2",
    "image-3",
    "image-4",
    "image-5",
    "image-6",
    "display_name",
    "product_description",
    "long_description",
    "section_description",
    "brand",
    "gtin",
    "barcode",
    "primary_image",
    "supplier_item_number",
    "primary_uom",
    "default_variant",
    "key_selling_point_1",
    "key_selling_point_2",
    "key_selling_point_3",
    "warranty_information",
})
# Listing fields that satisfy Bunnings operator attribute codes (PM11 / P41).
OPERATOR_FROM_LISTING = {
    "display_name": "title",
    "title": "title",
    "product_description": "description",
    "long_description": "description",
    "description": "description",
    "section_description": "description",
    "brand": "brand",
    "gtin": "gtin",
    "ean": "gtin",
    "barcode": "gtin",
    "category": "category",
}


def attributes_from_data(data: dict) -> dict:
    raw = data.get("attributes")
    if isinstance(raw, str) and raw.strip():
        try:
            raw = json.loads(raw)
        except (TypeError, ValueError, json.JSONDecodeError):
            raw = {}
    out = {}
    if isinstance(raw, dict):
        for key, value in raw.items():
            code = str(key or "").strip()
            if code and value not in (None, ""):
                out[code] = str(value).strip()
    for key, value in (data or {}).items():
        if key in ("attributes",):
            continue
        text = str(key or "").strip()
        if text.startswith("attribute_") or text.startswith("pdb_"):
            if value not in (None, ""):
                out[text] = str(value).strip()
    return out


def build_extras(data: dict) -> str:
    extras = {
        "marketplace": "bunnings",
        "gtin": str(data.get("gtin") or data.get("barcode") or "").strip(),
        "mpn": str(data.get("mpn") or "").strip(),
        "logistic_class": str(data.get("logistic_class") or "").strip(),
        "leadtime_to_ship": str(data.get("leadtime_to_ship") or "").strip() or "2",
        "product_id_type": str(data.get("product_id_type") or "").strip() or PRODUCT_ID_TYPE_SHOP_SKU,
        "weight": str(data.get("weight") or "").strip(),
        "weight_unit": str(data.get("weight_unit") or "").strip() or "kg",
        "length": str(data.get("length") or "").strip(),
        "height": str(data.get("height") or "").strip(),
        "width": str(data.get("width") or "").strip(),
        "dimension_unit": str(data.get("dimension_unit") or "").strip() or "cm",
        "attributes": attributes_from_data(data),
    }
    return json.dumps(extras)


def listing_sku(listing: StoreListing) -> str:
    return (listing.sku or listing.external_variant_key or "").strip()


def _slug_attr(name: str) -> str:
    text = str(name or "").strip().lower()
    if not text:
        return ""
    if text in OPTION_ATTR_ALIASES:
        return OPTION_ATTR_ALIASES[text]
    out: list[str] = []
    for ch in text:
        if ch.isalnum():
            out.append(ch)
        elif out and out[-1] != "-":
            out.append("-")
    return "".join(out).strip("-")


def collect_option_pairs(listing_or_data) -> list[tuple[str, str]]:
    data = listing_or_data
    if not isinstance(listing_or_data, dict):
        data = {
            f"option_{i}_{part}": getattr(listing_or_data, f"option_{i}_{part}", "")
            for i in (1, 2, 3, 4)
            for part in ("name", "value")
        }
    pairs: list[tuple[str, str]] = []
    for i in (1, 2, 3, 4):
        name = str(data.get(f"option_{i}_name") or "").strip()
        value = str(data.get(f"option_{i}_value") or "").strip()
        if name or value:
            pairs.append((name, value))
    return pairs


def variant_group_code(listing_or_data, sku: str = "") -> str:
    if isinstance(listing_or_data, dict):
        key = str(
            listing_or_data.get("product_key")
            or listing_or_data.get("variant_group_code")
            or ""
        ).strip()
        sku = sku or str(listing_or_data.get("sku") or listing_or_data.get("variant_key") or "").strip()
    else:
        key = str(getattr(listing_or_data, "external_product_key", "") or "").strip()
        sku = sku or listing_sku(listing_or_data)
    if key:
        return key
    if collect_option_pairs(listing_or_data):
        return sku
    return ""


def listing_price(listing: StoreListing) -> Decimal:
    price = _dec(listing.sale_price) if listing.sale_price else _dec(listing.original_price)
    if price > 0:
        return price
    cents = listing.sale_price_cents or listing.original_price_cents or 0
    return (Decimal(cents) / Decimal("100")) if cents else Decimal("0")


def _looks_like_placeholder(value) -> bool:
    return "REPLACE_WITH_" in str(value or "").upper()


def validate_listing(data: dict, store=None) -> list[str]:
    errors: list[str] = []
    sku = str(data.get("sku") or data.get("variant_key") or data.get("product_key") or "").strip()
    label = sku or "unknown"
    if not sku:
        errors.append("SKU is required for Bunnings.")
    if not str(data.get("title") or "").strip():
        errors.append(f"Title is required for SKU {label}.")
    if not str(data.get("description") or "").strip():
        errors.append(f"Description is required for SKU {label}.")
    if not str(data.get("brand") or "").strip():
        errors.append(f"Brand is required for SKU {label}.")
    elif _looks_like_placeholder(data.get("brand")):
        errors.append(
            f"Brand for SKU {label} is still a placeholder. "
            "Use a brand already approved on your Bunnings shop."
        )
    category = str(data.get("category") or data.get("category_uuid") or "").strip()
    if not category:
        errors.append(f"Bunnings category code is required for SKU {label}.")
    elif _looks_like_placeholder(category):
        errors.append(
            f"Category for SKU {label} is still a placeholder. "
            "Pick the hierarchy code from the Bunnings category search."
        )
    if not str(data.get("logistic_class") or "").strip():
        errors.append(
            f"Logistic class is required for SKU {label} "
            "(from Bunnings Mirakl shop settings)."
        )
    elif _looks_like_placeholder(data.get("logistic_class")):
        errors.append(
            f"Logistic class for SKU {label} is still a placeholder. "
            "Use a code from the Bunnings logistic class dropdown."
        )
    photos = _photo_urls(data.get("image_urls"))
    if not photos:
        errors.append(f"At least one image URL is required for SKU {label}.")
    elif any(_looks_like_placeholder(url) for url in photos):
        errors.append(
            f"Image URL for SKU {label} is still a placeholder. "
            "Use a public https image URL."
        )
    price = _dec(data.get("sale_price"))
    if price <= 0:
        price = _dec(data.get("original_price"))
    if price <= 0:
        errors.append(f"Price must be greater than 0 for SKU {label} (GST inclusive).")
    pairs = collect_option_pairs(data)
    for name, value in pairs:
        if not name or not value:
            errors.append(
                f"Option name and value must both be set for SKU {label} "
                "(e.g. Size / M)."
            )
            break
    if pairs:
        product_key = str(data.get("product_key") or "").strip()
        if not product_key:
            errors.append(
                f"Parent SKU is required for variation listings (SKU {label}). "
                "Use the same Parent SKU on every size/colour and a unique SKU per row."
            )
        elif product_key == sku:
            errors.append(
                f"Parent SKU must differ from SKU {label} so Bunnings can group "
                "sizes/colours on one product page."
            )
    attrs_in = attributes_from_data(data)
    category = str(data.get("category") or data.get("category_uuid") or "").strip()
    if store is not None and category and not _looks_like_placeholder(category):
        for item in load_category_attributes(store, category):
            if not item.get("required"):
                continue
            code = item.get("code") or ""
            if _attribute_satisfied(code, data, attrs_in):
                continue
            label_attr = item.get("label") or code
            errors.append(
                f"{label_attr} is required for this Bunnings category (SKU {label})."
            )
    return errors


def flatten_product_attributes(payload) -> list[dict]:
    """PM11 → form fields. Only required/recommended, skip core P41 columns."""
    items = []
    if isinstance(payload, dict):
        items = payload.get("attributes") or payload.get("data") or []
    elif isinstance(payload, list):
        items = payload
    if not isinstance(items, list):
        return []
    out = []
    seen = set()
    for raw in items:
        if not isinstance(raw, dict):
            continue
        code = str(raw.get("code") or raw.get("name") or "").strip()
        label = str(raw.get("label") or raw.get("description") or code).strip()
        attr_type = str(raw.get("type") or "TEXT").strip()
        extra_image = is_additional_image_attribute(code, label, attr_type)
        if not code or code in seen:
            continue
        if (code in CORE_PRODUCT_CODES or code.startswith("image-")) and not extra_image:
            continue
        seen.add(code)
        level = str(raw.get("requirement_level") or raw.get("requirementLevel") or "").strip().upper()
        required = bool(raw.get("required")) or level in ("REQUIRED", "MANDATORY")
        recommended = level in ("RECOMMENDED",)
        variant = bool(raw.get("variant"))
        if not required and not recommended and not variant and not extra_image:
            continue
        values = []
        raw_values = raw.get("values") or []
        if isinstance(raw_values, list):
            for val in raw_values:
                if isinstance(val, dict):
                    vcode = str(val.get("code") or val.get("label") or "").strip()
                    vlabel = str(val.get("label") or val.get("code") or vcode).strip()
                    if vcode:
                        values.append({"code": vcode, "label": vlabel})
                elif val not in (None, ""):
                    values.append({"code": str(val), "label": str(val)})
        out.append({
            "code": code,
            "label": label,
            "required": required,
            "recommended": recommended,
            "type": attr_type,
            "values": values,
            "values_list": _values_list_code(raw),
            "variant": variant,
            "extra_image": extra_image,
        })
    out.sort(key=lambda r: (not r["required"], r["label"].lower()))
    return out


def template_attribute_columns(store, hierarchy_codes) -> list[dict]:
    """PM11 columns for a bulk CSV. Header is the Bunnings attribute code."""
    seen: set[str] = set()
    columns: list[dict] = []
    for raw_code in hierarchy_codes or []:
        hierarchy = str(raw_code or "").strip()
        if not hierarchy:
            continue
        for item in load_category_attributes(store, hierarchy):
            code = str(item.get("code") or "").strip()
            if not code or code in seen or code.lower() in TEMPLATE_SKIP_ATTRS:
                continue
            if item.get("extra_image") or additional_image_slot(code):
                continue
            seen.add(code)
            columns.append({
                "code": code,
                "header": code,
                "label": str(item.get("label") or code).strip() or code,
                "required": bool(item.get("required")),
            })
    return columns


def load_category_attributes(store, hierarchy_code: str) -> list[dict]:
    code = str(hierarchy_code or "").strip()
    if not code or store is None:
        return []
    cache_key = (str(getattr(store, "id", "")), code)
    with _ATTR_LOCK:
        cached = _ATTR_CACHE.get(cache_key)
    if cached is not None:
        return cached
    try:
        client = BunningsClient(store)
        result = client.list_product_attributes(code)
    except Exception as exc:  # noqa: BLE001
        logger.warning("Bunnings PM11 failed hierarchy=%s: %s", code, exc)
        return []
    if not result.ok:
        logger.warning("Bunnings PM11 rejected hierarchy=%s: %s", code, result.message)
        return []
    rows = flatten_product_attributes(result.data)
    attach_list_values(store, rows)
    with _ATTR_LOCK:
        _ATTR_CACHE[cache_key] = rows
    return rows


def prefetch_category_attributes(store, hierarchy_codes, *, workers: int = 8) -> None:
    """Load PM11 for many hierarchies at once so a full-tree template can finish."""
    codes: list[str] = []
    seen: set[str] = set()
    for raw in hierarchy_codes or []:
        code = str(raw or "").strip()
        if code and code not in seen:
            seen.add(code)
            codes.append(code)
    if not codes or store is None:
        return
    load_value_lists(store)
    pool = max(1, min(int(workers or 1), len(codes), 8))
    if pool == 1:
        for code in codes:
            load_category_attributes(store, code)
        return
    from concurrent.futures import ThreadPoolExecutor

    with ThreadPoolExecutor(max_workers=pool) as executor:
        list(executor.map(lambda code: load_category_attributes(store, code), codes))


def _values_list_code(raw: dict) -> str:
    if not isinstance(raw, dict):
        return ""
    skip = {"", "LIST", "TEXT", "INTEGER", "DECIMAL", "BOOLEAN", "DATE", "MEDIA", "LONG_TEXT"}
    for key in ("values_list", "values_list_code", "type_parameter"):
        val = raw.get(key)
        if isinstance(val, dict):
            val = val.get("value") or val.get("code")
        code = str(val or "").strip()
        if code and code.upper() not in skip:
            return code
    for item in raw.get("type_parameters") or []:
        if not isinstance(item, dict):
            continue
        if str(item.get("name") or "").strip().upper() in ("LIST_CODE", "VALUES_LIST"):
            return str(item.get("value") or item.get("code") or "").strip()
    return ""


def load_value_lists(store) -> dict[str, list[dict]]:
    """code → [{code, label}, ...] from Mirakl VL11."""
    if store is None:
        return {}
    cache_key = str(getattr(store, "id", "") or "")
    if cache_key and cache_key in _VALUE_LIST_CACHE:
        return _VALUE_LIST_CACHE[cache_key]
    out: dict[str, list[dict]] = {}
    try:
        client = BunningsClient(store)
        result = client.get("/api/values_lists", params={"max": 10000})
    except Exception as exc:  # noqa: BLE001
        logger.warning("Bunnings value lists failed store=%s: %s", getattr(store, "id", None), exc)
        return {}
    data = getattr(result, "data", None)
    if getattr(result, "ok", False) is not True or not isinstance(data, dict):
        if getattr(result, "ok", None) is False:
            logger.warning(
                "Bunnings value lists rejected store=%s: %s",
                getattr(store, "id", None),
                getattr(result, "message", None),
            )
        return {}
    for lst in data.get("values_lists") or []:
        if not isinstance(lst, dict):
            continue
        list_code = str(lst.get("code") or "").strip()
        if not list_code:
            continue
        entries = []
        for val in lst.get("values") or []:
            if not isinstance(val, dict):
                continue
            vcode = str(val.get("code") or "").strip()
            if not vcode:
                continue
            entries.append({
                "code": vcode,
                "label": str(val.get("label") or vcode).strip() or vcode,
            })
        out[list_code] = entries
    if cache_key and out:
        _VALUE_LIST_CACHE[cache_key] = out
    return out


def attach_list_values(store, attributes: list[dict]) -> list[dict]:
    """Fill attribute.values from VL11 when PM11 left values empty."""
    if not attributes:
        return attributes
    lists = load_value_lists(store)
    if not lists:
        return attributes
    for attr in attributes:
        if not isinstance(attr, dict):
            continue
        if attr.get("values"):
            continue
        list_code = str(attr.get("values_list") or "").strip() or DEFAULT_ATTR_VALUE_LISTS.get(
            str(attr.get("code") or "").strip().upper(),
            "",
        )
        entries = lists.get(list_code) or []
        if entries and len(entries) <= MAX_INLINE_LIST_VALUES:
            attr["values"] = entries
    return attributes


def resolve_list_value(value_lists: dict | None, list_code: str, raw) -> str:
    """Map a typed label (Shemaya / Each) to the Mirakl list code (684010 / EA)."""
    text = str(raw or "").strip()
    if not text:
        return ""
    entries = (value_lists or {}).get(str(list_code or "").strip()) or []
    if not entries:
        return text
    for item in entries:
        if str(item.get("code") or "").strip() == text:
            return str(item["code"]).strip()
    low = text.lower()
    for item in entries:
        code = str(item.get("code") or "").strip()
        label = str(item.get("label") or "").strip()
        if code.lower() == low or label.lower() == low:
            return code
    return text


def _list_code_for_attr(attr_name: str, list_by_attr: dict | None) -> str:
    key = str(attr_name or "").strip()
    mapping = list_by_attr or {}
    return (
        mapping.get(key)
        or mapping.get(key.upper())
        or mapping.get(key.lower())
        or DEFAULT_ATTR_VALUE_LISTS.get(key.upper(), "")
    )


def _attr_lookup(attrs: dict, code: str) -> str:
    if not isinstance(attrs, dict):
        return ""
    if str(attrs.get(code) or "").strip():
        return str(attrs.get(code)).strip()
    low = (code or "").strip().lower()
    for key, value in attrs.items():
        if str(key or "").strip().lower() == low and value not in (None, ""):
            return str(value).strip()
    return ""


def _attribute_satisfied(code: str, data: dict, attrs: dict) -> bool:
    if _attr_lookup(attrs, code):
        return True
    low = (code or "").strip().lower().replace("-", "_")
    mapped = OPERATOR_FROM_LISTING.get(low)
    if mapped == "gtin":
        return bool(str(data.get("gtin") or data.get("barcode") or "").strip())
    if mapped and str(data.get(mapped) or "").strip():
        return True
    if low in ("primary_image",) or low.startswith("image"):
        return bool(_photo_urls(data.get("image_urls")))
    image_slot = additional_image_slot(code)
    if image_slot:
        photos = _photo_urls(data.get("image_urls"))
        if len(photos) > image_slot or _attr_lookup(attrs, code):
            return True
    extras_map = {
        "weight": "weight",
        "product-weight": "weight",
        "gross-weight": "weight",
        "length": "length",
        "product-length": "length",
        "package-length": "length",
        "height": "height",
        "product-height": "height",
        "width": "width",
        "product-width": "width",
    }
    dim = extras_map.get((code or "").strip().lower())
    if dim and str(data.get(dim) or "").strip():
        return True
    return False


def _write_csv(headers: list[str], rows: list[dict]) -> str:
    buf = io.StringIO()
    writer = csv.DictWriter(
        buf,
        fieldnames=headers,
        delimiter=";",
        quoting=csv.QUOTE_MINIMAL,
        lineterminator="\n",
        extrasaction="ignore",
    )
    writer.writeheader()
    for row in rows:
        writer.writerow({k: row.get(k, "") for k in headers})
    return buf.getvalue()


def _assign_additional_images(row: dict, gallery: list[str], codes: list[str], attrs: dict) -> None:
    """Write extra photo URLs onto the attribute codes Bunnings reads."""
    by_slot: dict[int, list[str]] = {}
    for code in codes:
        slot = additional_image_slot(code) or (len(by_slot) + 1)
        by_slot.setdefault(slot, []).append(code)
    for slot, slot_codes in by_slot.items():
        url = gallery[slot] if slot < len(gallery) else ""
        if not url:
            for code in slot_codes:
                url = _attr_lookup(attrs, code)
                if url:
                    break
        if not url:
            continue
        for code in slot_codes:
            row[code] = url


def product_row(
    listing: StoreListing,
    value_lists: dict | None = None,
    list_by_attr: dict | None = None,
    extra_image_codes: list[str] | None = None,
) -> dict:
    extras = parse_extras(listing)
    sku = listing_sku(listing)
    gtin = extras.get("gtin") or (listing.barcode or "").strip()
    title = (listing.title or sku)[:500]
    description = listing.description or listing.title or sku
    brand = (listing.brand or "").strip()
    category = (listing.category or "").strip()
    extras_attrs = extras.get("attributes") if isinstance(extras.get("attributes"), dict) else {}
    value_lists = value_lists or {}
    list_by_attr = list_by_attr or {}
    row = {
        "category": category,
        "product-id": sku,
        "product-id-type": extras.get("product_id_type") or PRODUCT_ID_TYPE_SHOP_SKU,
        VARIANT_GROUP_HEADER: "",
        OPERATOR_VARIANT_GROUP: "",
        "title": title,
        "description": description,
        "brand": brand,
        "ean": gtin,
        "DISPLAY_NAME": _attr_lookup(extras_attrs, "DISPLAY_NAME") or title,
        "PRODUCT_DESCRIPTION": _attr_lookup(extras_attrs, "PRODUCT_DESCRIPTION") or description,
        "LONG_DESCRIPTION": _attr_lookup(extras_attrs, "LONG_DESCRIPTION") or description,
        "BRAND": _attr_lookup(extras_attrs, "BRAND") or brand,
        "GTIN": _attr_lookup(extras_attrs, "GTIN") or gtin,
        "PRIMARY_IMAGE": _attr_lookup(extras_attrs, "PRIMARY_IMAGE") or "",
        "CATEGORY": _attr_lookup(extras_attrs, "CATEGORY") or category,
    }
    group = variant_group_code(listing, sku) or _attr_lookup(extras_attrs, OPERATOR_VARIANT_GROUP)
    row[VARIANT_GROUP_HEADER] = group
    row[OPERATOR_VARIANT_GROUP] = group
    gallery = gallery_urls(listing, extras_attrs)
    if gallery and not row["PRIMARY_IMAGE"]:
        row["PRIMARY_IMAGE"] = gallery[0]
    codes = extra_image_codes
    if codes is None:
        codes = extra_image_codes_for(getattr(listing, "store", None), [category])
    _assign_additional_images(row, gallery, codes, extras_attrs)
    for name, value in collect_option_pairs(listing):
        code = _slug_attr(name)
        if code and value:
            row[code] = value
    skip_extra = CORE_PRODUCT_CODES | {h.lower() for h in OFFER_CSV_HEADERS} | {
        "update-delete",
        "update_delete",
        "update delete",
    }
    for code, value in extras_attrs.items():
        key = str(code or "").strip()
        if not key or value in (None, ""):
            continue
        if key.lower() in skip_extra or additional_image_slot(key):
            continue
        if not str(row.get(key) or "").strip():
            row[key] = str(value).strip()
    for field, aliases in _DIM_FALLBACKS.items():
        val = str(extras.get(field) or "").strip()
        if not val:
            continue
        if not any(row.get(alias) for alias in aliases):
            row[aliases[0]] = val
    for key, value in list(row.items()):
        list_code = _list_code_for_attr(key, list_by_attr)
        if list_code:
            row[key] = resolve_list_value(value_lists, list_code, value)
    return row


def offer_row(listing: StoreListing, *, delete: bool = False) -> dict:
    extras = parse_extras(listing)
    sku = listing_sku(listing)
    qty = 0 if listing.infinite_quantity else max(0, int(listing.inventory or 0))
    if listing.infinite_quantity:
        qty = 9999
    return {
        "sku": sku,
        "product-id": sku,
        "product-id-type": extras.get("product_id_type") or PRODUCT_ID_TYPE_SHOP_SKU,
        "description": (listing.title or sku)[:2000],
        "price": f"{listing_price(listing):.2f}",
        "quantity": "0" if delete else str(qty),
        "state": OFFER_STATE_NEW,
        "logistic-class": extras.get("logistic_class") or "",
        "leadtime-to-ship": extras.get("leadtime_to_ship") or "2",
        "update-delete": "delete" if delete else "update",
    }


def products_csv(listings: list[StoreListing]) -> str:
    store = getattr(listings[0], "store", None) if listings else None
    value_lists = load_value_lists(store) if store is not None else {}
    list_by_attr = dict(DEFAULT_ATTR_VALUE_LISTS)
    categories = {
        str(getattr(listing, "category", None) or parse_extras(listing).get("category") or "").strip()
        for listing in listings
    }
    if store is not None:
        for category in categories:
            if not category:
                continue
            for item in load_category_attributes(store, category):
                list_code = str(item.get("values_list") or "").strip()
                attr = str(item.get("code") or "").strip()
                if list_code and attr:
                    list_by_attr[attr.upper()] = list_code
    image_codes = extra_image_codes_for(store, categories)
    rows = [
        product_row(l, value_lists=value_lists, list_by_attr=list_by_attr, extra_image_codes=image_codes)
        for l in listings
    ]
    extra: list[str] = []
    known = set(PRODUCT_CSV_HEADERS)
    for row in rows:
        for key in row:
            if key not in known and key not in extra:
                extra.append(key)
    return _write_csv(PRODUCT_CSV_HEADERS + extra, rows)


def offers_csv(listings: list[StoreListing], *, delete: bool = False) -> str:
    return _write_csv(OFFER_CSV_HEADERS, [offer_row(l, delete=delete) for l in listings])


def flatten_hierarchies(payload, *, q: str = "", limit: int = 2000) -> list[dict]:
    """Turn H11 payload into searchable {code, name, path} rows."""
    items = []
    if isinstance(payload, dict):
        items = payload.get("hierarchies") or payload.get("data") or []
    elif isinstance(payload, list):
        items = payload
    if not isinstance(items, list):
        return []

    by_code = {}
    for raw in items:
        if not isinstance(raw, dict):
            continue
        code = str(raw.get("code") or raw.get("hierarchy_code") or "").strip()
        if not code:
            continue
        label = str(raw.get("label") or raw.get("name") or code).strip()
        parent = str(raw.get("parent_code") or raw.get("parent") or "").strip()
        by_code[code] = {"code": code, "label": label, "parent_code": parent}

    def path_for(code: str) -> str:
        parts = []
        seen = set()
        cur = code
        while cur and cur not in seen:
            seen.add(cur)
            node = by_code.get(cur)
            if not node:
                break
            parts.append(node["label"])
            cur = node["parent_code"]
        return " / ".join(reversed(parts))

    children = {code: False for code in by_code}
    for node in by_code.values():
        parent = node["parent_code"]
        if parent in children:
            children[parent] = True

    query = (q or "").strip().lower()
    out = []
    for code, node in by_code.items():
        path = path_for(code)
        if query and query not in code.lower() and query not in path.lower():
            continue
        out.append({
            "code": code,
            "name": path or node["label"],
            "label": node["label"],
            "leaf": not children.get(code, False),
        })
    out.sort(key=lambda r: (not r["leaf"], r["name"].lower()))
    if limit and len(out) > limit:
        return out[:limit]
    return out


def flatten_logistic_classes(payload) -> list[dict]:
    items = []
    if isinstance(payload, dict):
        items = (
            payload.get("logistic_classes")
            or payload.get("logistics")
            or payload.get("data")
            or []
        )
    elif isinstance(payload, list):
        items = payload
    if not isinstance(items, list):
        return []
    out = []
    for raw in items:
        if not isinstance(raw, dict):
            continue
        code = str(raw.get("code") or raw.get("logistic_class") or "").strip()
        if not code:
            continue
        label = str(raw.get("label") or raw.get("name") or code).strip()
        out.append({"code": code, "name": label})
    return out


def _catalog_product_known(client, sku: str):
    """True/False when the client can answer; None skips the wait (tests / mocks)."""
    fn = getattr(client, "product_exists", None)
    if not callable(fn):
        return None
    try:
        found = fn(sku)
    except Exception:  # noqa: BLE001
        return None
    if isinstance(found, bool):
        return found
    return None


def wait_for_catalog_products(client, skus, *, attempts: int | None = None, interval: float | None = None) -> list[str]:
    """Poll until shop product-ids exist. Returns SKUs still missing."""
    remaining = [str(sku or "").strip() for sku in skus if str(sku or "").strip()]
    tries = attempts if attempts is not None else PRODUCT_WAIT_ATTEMPTS
    wait = interval if interval is not None else PRODUCT_WAIT_SECONDS
    for i in range(max(1, tries)):
        still = []
        unknown = False
        for sku in remaining:
            known = _catalog_product_known(client, sku)
            if known is None:
                unknown = True
                break
            if not known:
                still.append(sku)
        if unknown:
            return []
        remaining = still
        if not remaining:
            return []
        if i < tries - 1 and wait:
            time.sleep(wait)
    return remaining


def _offer_product_missing(result) -> bool:
    text = f"{getattr(result, 'message', None) or ''}".lower()
    if any(marker in text for marker in _PRODUCT_MISSING_MARKERS):
        return True
    data = getattr(result, "data", None)
    if not isinstance(data, dict):
        return False
    for row in data.get("line_errors") or []:
        if not isinstance(row, dict):
            continue
        err = str(row.get("errors") or "").lower()
        if any(marker in err for marker in _PRODUCT_MISSING_MARKERS):
            return True
    return False


def import_offers_waiting_for_products(client, listings, *, attempts: int | None = None, interval: float | None = None):
    """OF01 + poll. Retry when Mirakl says the product is not in the catalog yet."""
    tries = attempts if attempts is not None else OFFER_RETRY_ATTEMPTS
    wait = interval if interval is not None else OFFER_RETRY_SECONDS
    last = None
    for i in range(max(1, tries)):
        offer_text = offers_csv(listings)
        last = client.import_offers(offer_text)
        import_id = extract_import_id(getattr(last, "data", None))
        if last.ok and import_id:
            last = client.poll_import("offer", import_id)
        last.offer_csv = offer_text
        if last.ok or not _offer_product_missing(last):
            return last
        if i < tries - 1 and wait:
            time.sleep(wait)
    return last


_FIELD_LABELS = {
    "brand": "Brand",
    "category": "Category",
    "primary_uom": "Primary unit of measure",
    "gtin": "Barcode",
    "ean": "Barcode",
    "barcode": "Barcode",
    "display_name": "Website Display Name",
    "title": "Title",
    "description": "Description",
    "long_description": "Product Overview",
    "product_description": "Product Description",
    "section_description": "Section Description",
    "logistic_class": "Logistic class",
    "leadtime_to_ship": "Lead time to ship",
    "price": "Offer price",
    "quantity": "Offer quantity",
    "image_1": "Image",
    "primary_image": "Image",
    "state": "Offer state",
    "product_id": "Product ID",
    "sku": "SKU",
    "supplier_item_number": "Supplier item number",
}


def _norm_attr(code: str) -> str:
    return "".join(ch if ch.isalnum() else "_" for ch in str(code or "").strip().lower()).strip("_")


def _field_label(attribute: str) -> str:
    norm = _norm_attr(attribute)
    if norm in _FIELD_LABELS:
        return _FIELD_LABELS[norm]
    if not str(attribute or "").strip():
        return ""
    return str(attribute).replace("_", " ").replace("-", " ").strip().title()


def _add_sent_value(bag: dict[str, set[str]], code, value) -> None:
    key = _norm_attr(code)
    text = str(value or "").strip()
    if key and text:
        bag.setdefault(key, set()).add(text.casefold())


def _listing_sent_values(listing: StoreListing) -> dict[str, set[str]]:
    """Values we sent, so a Mirakl row with no SKU can still be tied to this product."""
    bag: dict[str, set[str]] = {}
    _add_sent_value(bag, "sku", listing_sku(listing))
    _add_sent_value(bag, "product-id", listing_sku(listing))
    _add_sent_value(bag, "brand", getattr(listing, "brand", ""))
    _add_sent_value(bag, "category", getattr(listing, "category", ""))
    _add_sent_value(bag, "title", getattr(listing, "title", ""))
    _add_sent_value(bag, "description", getattr(listing, "description", ""))
    _add_sent_value(bag, "gtin", getattr(listing, "barcode", ""))
    _add_sent_value(bag, "ean", getattr(listing, "barcode", ""))
    extras = parse_extras(listing)
    for key in ("gtin", "mpn", "logistic_class", "leadtime_to_ship"):
        _add_sent_value(bag, key, extras.get(key))
    attrs = extras.get("attributes") if isinstance(extras.get("attributes"), dict) else {}
    for code, value in attrs.items():
        _add_sent_value(bag, code, value)
    try:
        for code, value in product_row(listing).items():
            _add_sent_value(bag, code, value)
        for code, value in offer_row(listing).items():
            _add_sent_value(bag, code, value)
    except Exception as exc:  # noqa: BLE001
        logger.warning("Bunnings sent-value lookup failed for %s: %s", listing_sku(listing), exc)
    return bag


def _assign_import_errors(listings: list[StoreListing], rows: list[dict]) -> dict[int, list[dict]]:
    sent = {id(listing): _listing_sent_values(listing) for listing in listings}
    by_sku: dict[str, list[StoreListing]] = {}
    for listing in listings:
        for key in (listing_sku(listing), getattr(listing, "external_variant_key", "") or ""):
            text = str(key or "").strip()
            if text:
                by_sku.setdefault(text.casefold(), []).append(listing)
    assigned: dict[int, list[dict]] = {id(listing): [] for listing in listings}

    def targets(row: dict) -> list[StoreListing]:
        sku = str(row.get("sku") or "").strip()
        if sku:
            return list(by_sku.get(sku.casefold(), []))
        attr = _norm_attr(row.get("attribute"))
        val = str(row.get("value") or "").strip()
        if attr and val:
            hits = [
                listing for listing in listings
                if val.casefold() in (sent[id(listing)].get(attr) or set())
            ]
            if hits:
                return hits
            return list(listings) if len(listings) == 1 else []
        if len(listings) == 1:
            return list(listings)
        text = f"{row.get('errors') or ''} {row.get('message') or ''}".casefold()
        return [
            listing for listing in listings
            if listing_sku(listing) and listing_sku(listing).casefold() in text
        ]

    for row in rows:
        if not isinstance(row, dict):
            continue
        for listing in targets(row):
            assigned[id(listing)].append(row)
    return assigned


def _public_bunnings_error(listing: StoreListing, row: dict) -> dict:
    attribute = str(row.get("attribute") or "").strip()
    message = str(row.get("message") or "").strip()
    errors = str(row.get("errors") or "").strip()
    if not message:
        message = errors
    return {
        "sku": listing_sku(listing) or str(row.get("sku") or "").strip(),
        "code": str(row.get("code") or "").strip(),
        "attribute": attribute,
        "field": _field_label(attribute),
        "value": str(row.get("value") or "").strip(),
        "message": message,
        "errors": errors or message,
    }


def fail_listings_from_import(listings: list[StoreListing], result, *, request: dict) -> None:
    """Mark each product not created and keep the Mirakl error rows that belong to it."""
    data = getattr(result, "data", None)
    rows = line_error_rows(data)
    assigned = _assign_import_errors(listings, rows)
    fallback = (getattr(result, "message", None) or "").strip() or "Bunnings did not create this product."
    safe_data = data if isinstance(data, (dict, list)) else {"raw": "" if data is None else str(data)[:4000]}
    for listing in listings:
        public = []
        messages = []
        seen: set[tuple] = set()
        for row in assigned.get(id(listing)) or []:
            item = _public_bunnings_error(listing, row)
            key = (
                item["code"],
                item["attribute"].casefold(),
                item["value"].casefold(),
                item["message"],
            )
            if key in seen:
                continue
            seen.add(key)
            public.append(item)
            text = item["errors"] or item["message"]
            if text and text not in messages:
                messages.append(text)
        if not messages:
            messages = [fallback]
            public = []
        _mark_listing(
            listing,
            status=ListingStatus.FAILED,
            request=request,
            response={
                "error": messages[0],
                "data": safe_data,
                "bunnings_errors": public,
            },
            errors=messages,
        )


def stored_bunnings_errors(listing: StoreListing) -> list[dict]:
    raw = getattr(listing, "marketplace_response_json", None)
    rows = raw.get("bunnings_errors") if isinstance(raw, dict) else None
    if not isinstance(rows, list):
        return []
    return [row for row in rows if isinstance(row, dict)]


def error_report_rows(listings: list[StoreListing]) -> list[dict]:
    """Flat rows for the Excel file a seller downloads from Created products."""
    out: list[dict] = []
    for listing in listings:
        public = stored_bunnings_errors(listing)
        if not public:
            texts = getattr(listing, "validation_errors_json", None) or []
            if isinstance(texts, str):
                texts = [texts]
            public = [
                {
                    "sku": listing_sku(listing),
                    "field": "",
                    "value": "",
                    "code": "",
                    "message": str(text).strip(),
                }
                for text in texts
                if str(text or "").strip()
            ]
        if not public:
            public = [{
                "sku": listing_sku(listing),
                "field": "",
                "value": "",
                "code": "",
                "message": "Bunnings did not create this product.",
            }]
        for item in public:
            out.append({
                "sku": item.get("sku") or listing_sku(listing),
                "field": item.get("field") or _field_label(item.get("attribute") or ""),
                "value": item.get("value") or "",
                "code": item.get("code") or "",
                "message": item.get("message") or item.get("errors") or "",
                "result": "Not created",
            })
    return out


def error_report_xlsx(listings: list[StoreListing]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Not created"
    headers = ["SKU", "Field", "Value sent", "Error code", "Message", "Result"]
    bold = Font(bold=True)
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(1, col, header)
        cell.font = bold
    for row_i, item in enumerate(error_report_rows(listings), start=2):
        sheet.cell(row_i, 1, item["sku"])
        sheet.cell(row_i, 2, item["field"])
        sheet.cell(row_i, 3, item["value"])
        sheet.cell(row_i, 4, item["code"])
        sheet.cell(row_i, 5, item["message"])
        sheet.cell(row_i, 6, item["result"])
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


def _mark_listing(listing: StoreListing, *, status: str, request: dict, response, errors=None):
    listing.status = status
    listing.marketplace_request_json = request
    listing.marketplace_response_json = response if isinstance(response, (dict, list)) else {"raw": response}
    listing.last_uploaded_at = timezone.now() if status.startswith("uploaded") else listing.last_uploaded_at
    listing.validation_errors_json = errors
    listing.save(
        update_fields=[
            "status",
            "marketplace_request_json",
            "marketplace_response_json",
            "last_uploaded_at",
            "validation_errors_json",
            "updated_at",
        ]
    )


def _partition_by_catalog(client, listings: list[StoreListing]) -> tuple[list, list]:
    """Split listings whose shop SKU is already in Bunnings vs those that still need P41."""
    found: list[StoreListing] = []
    missing: list[StoreListing] = []
    for listing in listings:
        if _catalog_product_known(client, listing_sku(listing)) is True:
            found.append(listing)
        else:
            missing.append(listing)
    return found, missing


def publish_listings(user, store, listings: list[StoreListing]) -> dict:
    """P41 product import (create only) then OF01 offer import.

    If Bunnings already has the shop SKU (Pending/accepted), skip P41 and
    send the offer so price and stock are not blocked by a SENT product import.
    """
    client = BunningsClient(store)
    creates = [l for l in listings if (l.action or ListingAction.CREATE) != ListingAction.MAPPED]
    mapped = [l for l in listings if (l.action or "") == ListingAction.MAPPED]
    target_status = (
        ListingStatus.UPLOADED_PRODUCTION
        if client.environment == "production"
        else ListingStatus.UPLOADED_STAGING
    )

    product_import_id = ""
    product_result = None
    p41_failed_ids: set = set()
    already_in_catalog, need_p41 = _partition_by_catalog(client, creates)
    if need_p41:
        csv_text = products_csv(need_p41)
        product_result = client.import_products(csv_text)
        product_import_id = extract_import_id(product_result.data)
        if not product_result.ok:
            for listing in need_p41:
                p41_failed_ids.add(listing.id)
            fail_listings_from_import(
                need_p41,
                product_result,
                request={"p41": csv_text[:4000]},
            )
            if not mapped and not already_in_catalog:
                return {
                    "ok": False,
                    "published": 0,
                    "failed": len(need_p41),
                    "message": product_result.message or "Bunnings product import failed.",
                    "environment": client.environment,
                }
        elif product_import_id:
            polled = client.poll_import("product", product_import_id)
            product_result = polled
            if _import_still_running(polled):
                now_found, _still_missing = _partition_by_catalog(client, need_p41)
                already_in_catalog = already_in_catalog + now_found
            elif not polled.ok:
                for listing in need_p41:
                    p41_failed_ids.add(listing.id)
                fail_listings_from_import(
                    need_p41,
                    polled,
                    request={"p41_import_id": product_import_id},
                )
                if not mapped and not already_in_catalog:
                    return {
                        "ok": False,
                        "published": 0,
                        "failed": len(need_p41),
                        "message": polled.message or "Bunnings product import failed.",
                        "environment": client.environment,
                    }
        if product_result and product_result.ok:
            wait_for_catalog_products(client, [listing_sku(l) for l in need_p41])

    # Previous Push failed rows must still get OF01 after a successful P41 retry.
    offer_targets = [l for l in listings if l.id not in p41_failed_ids]
    if not offer_targets:
        detail = ""
        if product_result is not None:
            detail = (product_result.message or "").strip()
        return {
            "ok": False,
            "published": 0,
            "failed": len(listings),
            "message": detail or "Bunnings product import failed; offer was not sent.",
            "environment": client.environment,
        }

    offer_result = import_offers_waiting_for_products(client, offer_targets)
    offer_import_id = extract_import_id(getattr(offer_result, "data", None))
    offer_text = getattr(offer_result, "offer_csv", "") or offers_csv(offer_targets)

    create_ids = {id(l) for l in creates}
    p41_accepted = bool(
        (product_result and (product_result.ok or _import_still_running(product_result)))
        or already_in_catalog
    )
    if (
        not offer_result.ok
        and _offer_product_missing(offer_result)
        and p41_accepted
    ):
        waiting = [l for l in offer_targets if id(l) in create_ids]
        hard_fail = [l for l in offer_targets if id(l) not in create_ids]
        for listing in waiting:
            _mark_listing(
                listing,
                status=ListingStatus.READY,
                request={"p41_import_id": product_import_id, "of01_import_id": offer_import_id},
                response={
                    "product_import": None if product_result is None else product_result.data,
                    "offer_import": offer_result.data,
                    "waiting_for_product": True,
                },
                errors=None,
            )
        if hard_fail:
            fail_listings_from_import(
                hard_fail,
                offer_result,
                request={"p41_import_id": product_import_id, "of01": offer_text[:4000]},
            )
        if waiting and not hard_fail:
            return {
                "ok": True,
                "published": 0,
                "failed": 0,
                "pending": len(waiting),
                "message": PENDING_OFFER_MESSAGE,
                "environment": client.environment,
            }
        return {
            "ok": False,
            "published": 0,
            "failed": len(hard_fail),
            "pending": len(waiting),
            "message": offer_result.message or PENDING_OFFER_MESSAGE,
            "environment": client.environment,
        }

    published = 0
    failed = 0
    if offer_result.ok:
        for listing in offer_targets:
            _mark_listing(
                listing,
                status=target_status,
                request={"p41_import_id": product_import_id, "of01_import_id": offer_import_id},
                response={
                    "product_import": None if product_result is None else product_result.data,
                    "offer_import": offer_result.data,
                },
            )
            published += 1
    else:
        fail_listings_from_import(
            offer_targets,
            offer_result,
            request={"p41_import_id": product_import_id, "of01": offer_text[:4000]},
        )
        failed = len(offer_targets)

    if published and failed == 0:
        message = (
            f"Published {published} listing(s) to Bunnings. "
            "New products may wait for Bunnings review before they are live."
        )
    elif published:
        message = f"Published {published} listing(s); {failed} failed."
    else:
        message = offer_result.message or "Could not publish listings to Bunnings."
    return {
        "ok": published > 0 and failed == 0,
        "published": published,
        "failed": failed,
        "message": message,
        "environment": client.environment,
    }


def _import_still_running(result) -> bool:
    msg = (result.message or "").lower()
    return "still" in msg or "in progress" in msg or "waiting" in msg or "running" in msg


def push_inventory(listings: list[StoreListing], store) -> dict:
    """OF01 offer update for price + stock."""
    if not listings:
        return {"ok": False, "message": "No listings to push.", "updated": 0}
    client = BunningsClient(store)
    csv_text = offers_csv(listings)
    result = client.import_offers(csv_text, import_mode="NORMAL")
    import_id = extract_import_id(result.data)
    if result.ok and import_id:
        polled = client.poll_import("offer", import_id)
        if polled.ok or _import_still_running(polled):
            result = polled if polled.ok else result
        else:
            result = polled
    return {
        "ok": result.ok,
        "message": result.message or ("Price/stock pushed to Bunnings." if result.ok else "Bunnings offer update failed."),
        "updated": len(listings) if result.ok else 0,
    }


def end_listing(store, listing: StoreListing) -> bool:
    sku = listing_sku(listing)
    if not sku:
        return False
    client = BunningsClient(store)
    csv_text = offers_csv([listing], delete=True)
    result = client.import_offers(csv_text, import_mode="NORMAL")
    if result.ok:
        return True
    msg = (result.message or "").lower()
    if result.status == 404 or "not found" in msg:
        return False
    raise MarketplaceError(result.message or "Could not end the Bunnings offer.")


def lookup_offer(store, sku: str) -> dict | None:
    client = BunningsClient(store)
    result = client.list_offers(sku=sku, max_results=20)
    if not result.ok:
        raise MarketplaceError(result.message or "Bunnings offer lookup failed.")
    if not isinstance(result.data, dict):
        return None
    offers = result.data.get("offers") or result.data.get("data") or []
    if not isinstance(offers, list):
        return None
    needle = sku.strip().lower()
    for offer in offers:
        if not isinstance(offer, dict):
            continue
        shop_sku = str(offer.get("shop_sku") or offer.get("sku") or "").strip()
        product_id = str(offer.get("product_id") or offer.get("product-id") or "").strip()
        if shop_sku.lower() == needle or product_id.lower() == needle:
            return offer
    return offers[0] if offers else None
