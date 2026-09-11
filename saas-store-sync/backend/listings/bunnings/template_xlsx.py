"""Bunnings bulk listing Excel template (operator-style colors + dropdowns)."""
from __future__ import annotations

import logging
from dataclasses import dataclass, field
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import products as bunnings_products
from .client import BunningsClient

logger = logging.getLogger("listings.bunnings")

COLOR_SYSTEM = "E67E22"
COLOR_PRODUCT = "70BB32"
COLOR_CATEGORY = "4F8A22"
COLOR_OFFER = "1563A3"
FONT_WHITE = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
ALIGN = Alignment(wrap_text=True, vertical="center", horizontal="center")
DATA_ROWS_END = 10000
YES_NO = ["Yes", "No"]
PRODUCT_ID_TYPES = ["SHOP_SKU", "GTIN", "SKU"]
ORIGIN_DEFAULTS = ["Australia", "New Zealand"]

BUNNINGS_PRODUCT_REQUIRED = [
    ("CATEGORY", "Category", "The product category (hierarchy code)."),
    ("SUPPLIER_ITEM_NUMBER", "Supplier Item Number", "Provide your product SKU."),
    ("DISPLAY_NAME", "Website Display Name", "Brand + item description + one unique feature."),
    ("BRAND", "Brand", "Brand already approved on your Bunnings shop."),
    ("GTIN", "Barcode", "One GTIN per SKU."),
    ("LONG_DESCRIPTION", "Product Overview", "Product overview paragraph. Not the same as key selling points."),
    ("SECTION_DESCRIPTION", "Section Description", "Maximum 30 characters. No special characters."),
    ("PRODUCT_DESCRIPTION", "Product Description", "Maximum 30 characters. No special characters."),
    ("PRIMARY_UOM", "Primary unit of measure", "Primary unit of measure."),
    ("PRIMARY_IMAGE", "Image", "Primary image, product only, white background."),
    ("VARIANT_GROUP_CODE", "Variant group code", "Same code on every size/colour variant."),
    ("DEFAULT_VARIANT", "Default Variant", "Yes on one row per variant group, No on the others."),
    ("KEY_SELLING_POINT_1", "Key Selling Point 1", "Short sentence, different from the product overview."),
    ("KEY_SELLING_POINT_2", "Key Selling Point 2", "Short sentence, different from the product overview."),
    ("KEY_SELLING_POINT_3", "Key Selling Point 3", "Short sentence, different from the product overview."),
    ("WARRANTY_INFORMATION", "Warranty Information", "Timeframe only, e.g. 12 Months."),
]
BUNNINGS_OFFER_REQUIRED = [
    ("sku", "Offer SKU", "Unique offer SKU. Also sent as product-id when that column matches."),
    ("product-id", "Product ID", "Unique product identifier for the product-id-type."),
    ("product-id-type", "Product ID Type", "SHOP_SKU, GTIN, or SKU."),
    ("price", "Offer Price", "GST inclusive price."),
    ("quantity", "Offer Quantity", "Quantity available in stock."),
    ("state", "Offer State", "Offer condition. Use New."),
    ("logistic-class", "Logistic Class", "Mirakl logistic class code from shop settings."),
    ("leadtime-to-ship", "Leadtime To Ship", "Days to ship. Defaults to 2."),
]
SYSTEM_COLUMNS = [
    ("Vendor Name (Optional)", "Vendor Name (Optional)", "Source vendor name for scrape routing."),
    ("Vendor URL (Optional)", "Vendor URL (Optional)", "Source product URL."),
    ("Vendor ID (Optional)", "Vendor ID (Optional)", "Source vendor product id."),
    ("Marketplace Name (Optional)", "Marketplace Name (Optional)", "Must match this store marketplace."),
    ("Store Name (Optional)", "Store Name (Optional)", "Must match this store name."),
    ("Action", "Action", "Create or Mapped."),
    ("Option 1 Name (Optional)", "Option 1 Name (Optional)", "e.g. Size."),
    ("Option 1 Value (Optional)", "Option 1 Value (Optional)", "e.g. M."),
    ("Option 2 Name (Optional)", "Option 2 Name (Optional)", "e.g. Colour."),
    ("Option 2 Value (Optional)", "Option 2 Value (Optional)", "e.g. Red."),
    ("Option 3 Name (Optional)", "Option 3 Name (Optional)", ""),
    ("Option 3 Value (Optional)", "Option 3 Value (Optional)", ""),
    ("Option 4 Name (Optional)", "Option 4 Name (Optional)", ""),
    ("Option 4 Value (Optional)", "Option 4 Value (Optional)", ""),
    ("Variation Img URL (Optional)", "Variation Img URL (Optional)", "Image for this size/colour."),
]
ALWAYS_PRODUCT_CODES = frozenset(code.lower() for code, _label, _prompt in BUNNINGS_PRODUCT_REQUIRED)
ALWAYS_OFFER_CODES = frozenset(code.lower() for code, _label, _prompt in BUNNINGS_OFFER_REQUIRED)


@dataclass
class TemplateColumn:
    code: str
    label: str
    band: str
    prompt: str = ""
    values: list[str] = field(default_factory=list)


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _band_color(band: str) -> str:
    return {
        "system": COLOR_SYSTEM,
        "product": COLOR_PRODUCT,
        "category": COLOR_CATEGORY,
        "offer": COLOR_OFFER,
    }.get(band, COLOR_PRODUCT)


def _unique_values(*groups) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for group in groups:
        for raw in group or []:
            if isinstance(raw, dict):
                val = str(raw.get("code") or raw.get("label") or "").strip()
            else:
                val = str(raw or "").strip()
            key = val.lower()
            if val and key not in seen:
                seen.add(key)
                out.append(val)
    return out


def _pm11_lookup(store, hierarchy_codes) -> dict[str, dict]:
    by_code: dict[str, dict] = {}
    for raw in hierarchy_codes or []:
        for item in bunnings_products.load_category_attributes(store, raw):
            code = str(item.get("code") or "").strip()
            if not code:
                continue
            prev = by_code.get(code.lower())
            values = _unique_values((prev or {}).get("values"), item.get("values"))
            required = bool(item.get("required")) or bool((prev or {}).get("required"))
            by_code[code.lower()] = {
                "code": code,
                "label": str(item.get("label") or (prev or {}).get("label") or code),
                "required": required,
                "values": values,
            }
    return by_code


def _logistic_codes(store) -> list[str]:
    if store is None:
        return []
    try:
        client = BunningsClient(store)
        result = client.list_logistic_classes()
    except Exception as exc:  # noqa: BLE001
        logger.warning("Bunnings logistic classes for template failed: %s", exc)
        return []
    if not result.ok:
        return []
    return [row["code"] for row in bunnings_products.flatten_logistic_classes(result.data) if row.get("code")]


def _dropdown_defaults(code: str, values: list[str]) -> list[str]:
    low = (code or "").strip().lower().replace("-", "_")
    if low == "default_variant" or low == "tax_inc":
        return _unique_values(values, YES_NO)
    if low == "primary_uom":
        return _unique_values(values, ["Each"])
    if low == "country_of_origin":
        return _unique_values(values, ORIGIN_DEFAULTS)
    if low == "category":
        return values
    return values


def build_template_columns(store, hierarchy_codes) -> list[TemplateColumn]:
    pm11 = _pm11_lookup(store, hierarchy_codes)
    categories = [str(c).strip() for c in (hierarchy_codes or []) if str(c or "").strip()]
    logistics = _logistic_codes(store)
    columns: list[TemplateColumn] = []

    for code, label, prompt in SYSTEM_COLUMNS:
        values = ["Create", "Mapped"] if code == "Action" else []
        columns.append(TemplateColumn(code=code, label=label, band="system", prompt=prompt, values=values))

    for code, label, prompt in BUNNINGS_PRODUCT_REQUIRED:
        extra = pm11.get(code.lower(), {})
        values = list(extra.get("values") or [])
        if code == "CATEGORY":
            values = categories
        else:
            values = _dropdown_defaults(code, values)
        columns.append(TemplateColumn(code=code, label=label, band="product", prompt=prompt, values=values))

    seen = {col.code.lower() for col in columns}
    for item in pm11.values():
        code = item["code"]
        if not item.get("required"):
            continue
        if code.lower() in seen or code.lower() in ALWAYS_PRODUCT_CODES or code.lower() in ALWAYS_OFFER_CODES:
            continue
        if code.lower() in bunnings_products.TEMPLATE_SKIP_ATTRS:
            continue
        seen.add(code.lower())
        columns.append(TemplateColumn(
            code=code,
            label=item["label"] or code,
            band="category",
            prompt=f"Required for the selected Bunnings categor{'y' if len(categories) == 1 else 'ies'}.",
            values=_dropdown_defaults(code, list(item.get("values") or [])),
        ))

    for code, label, prompt in BUNNINGS_OFFER_REQUIRED:
        values: list[str] = []
        if code == "product-id-type":
            values = list(PRODUCT_ID_TYPES)
        elif code == "state":
            values = ["New"]
        elif code == "logistic-class":
            values = logistics
        columns.append(TemplateColumn(code=code, label=label, band="offer", prompt=prompt, values=values))
    return columns


def _sample_row(columns: list[TemplateColumn], *, category: str, index: int, store_name: str, marketplace_name: str, action: str) -> dict:
    suffix = "" if index == 1 else f"-{index}"
    sku = f"BN-EXAMPLE-001{suffix}-M"
    parent = f"BN-EXAMPLE-001{suffix}"
    title = "Example Power Drill"
    overview = "Example product overview for Bunnings Marketplace. Keep key selling points separate."
    image = "https://example.com/photo1.jpg"
    logistic = ""
    uom = "Each"
    for col in columns:
        if col.code == "logistic-class" and col.values:
            logistic = col.values[0]
        if col.code == "PRIMARY_UOM" and col.values:
            uom = col.values[0]
    samples = {
        "Vendor Name (Optional)": "Amazon AU",
        "Vendor URL (Optional)": "https://www.amazon.com.au/dp/EXAMPLE",
        "Marketplace Name (Optional)": marketplace_name or "Bunnings",
        "Store Name (Optional)": store_name,
        "Action": "Mapped" if action == "mapped" else "Create",
        "Option 1 Name (Optional)": "Size",
        "Option 1 Value (Optional)": "M",
        "Option 2 Name (Optional)": "Colour",
        "Option 2 Value (Optional)": "Red",
        "Variation Img URL (Optional)": "https://example.com/photo-red-m.jpg",
        "CATEGORY": category,
        "SUPPLIER_ITEM_NUMBER": sku,
        "DISPLAY_NAME": title,
        "BRAND": "ExampleBrand",
        "GTIN": "9300000000001",
        "LONG_DESCRIPTION": overview,
        "SECTION_DESCRIPTION": "Example section text",
        "PRODUCT_DESCRIPTION": "Example product text",
        "PRIMARY_UOM": uom,
        "PRIMARY_IMAGE": image,
        "VARIANT_GROUP_CODE": parent,
        "DEFAULT_VARIANT": "Yes" if index == 1 else "No",
        "KEY_SELLING_POINT_1": "Example selling point 1",
        "KEY_SELLING_POINT_2": "Example selling point 2",
        "KEY_SELLING_POINT_3": "Example selling point 3",
        "WARRANTY_INFORMATION": "12 Months",
        "sku": sku,
        "product-id": sku,
        "product-id-type": "SHOP_SKU",
        "price": "79.99",
        "quantity": "5",
        "state": "New",
        "logistic-class": logistic,
        "leadtime-to-ship": "2",
    }
    return {col.code: samples.get(col.code, "") for col in columns}


def build_template_xlsx(action: str = "create", store=None, hierarchies=None) -> bytes:
    """Colored Bunnings listing workbook: system, product required, category required, offer."""
    action = (action or "create").strip().lower()
    hierarchy_codes = [str(c).strip() for c in (hierarchies or []) if str(c or "").strip()]
    columns = build_template_columns(store, hierarchy_codes)
    store_name = (getattr(store, "name", None) or "").strip()
    marketplace_name = ""
    mp = getattr(store, "marketplace", None)
    if mp is not None:
        marketplace_name = (getattr(mp, "name", None) or getattr(mp, "code", None) or "").strip()

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ref = wb.create_sheet("ReferenceData")

    fills = {band: _fill(_band_color(band)) for band in ("system", "product", "category", "offer")}
    for idx, col in enumerate(columns, start=1):
        fill = fills[col.band]
        label_cell = ws.cell(1, idx, col.label)
        code_cell = ws.cell(2, idx, col.code)
        for cell in (label_cell, code_cell):
            cell.fill = fill
            cell.font = FONT_WHITE
            cell.alignment = ALIGN
        ws.column_dimensions[get_column_letter(idx)].width = min(28, max(16, len(col.label) + 2))

    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 22
    ws.freeze_panes = "A3"
    last_col = get_column_letter(len(columns))
    ws.auto_filter.ref = f"A2:{last_col}2"

    samples = hierarchy_codes or [""]
    for i, code in enumerate(samples, start=1):
        values = _sample_row(
            columns,
            category=code or "HIERARCHY_CODE",
            index=i,
            store_name=store_name,
            marketplace_name=marketplace_name,
            action=action,
        )
        for idx, col in enumerate(columns, start=1):
            ws.cell(2 + i, idx, values.get(col.code, ""))

    ref_col = 1
    for col_idx, col in enumerate(columns, start=1):
        if not col.values:
            continue
        ref.cell(1, ref_col, col.code)
        for row_i, value in enumerate(col.values, start=2):
            ref.cell(row_i, ref_col, value)
        last = max(2, 1 + len(col.values))
        letter = get_column_letter(ref_col)
        formula = f"ReferenceData!${letter}$2:${letter}${last}"
        dv = DataValidation(
            type="list",
            formula1=formula,
            allow_blank=True,
            showDropDown=False,
            showInputMessage=True,
            showErrorMessage=False,
            promptTitle=(col.label or col.code)[:32],
            prompt=(col.prompt or "")[:255],
        )
        data_letter = get_column_letter(col_idx)
        dv.add(f"{data_letter}3:{data_letter}{DATA_ROWS_END}")
        ws.add_data_validation(dv)
        ref_col += 1

    for col_idx, col in enumerate(columns, start=1):
        if col.values or not col.prompt:
            continue
        dv = DataValidation(
            type="textLength",
            operator="lessThanOrEqual",
            formula1="32000",
            allow_blank=True,
            showInputMessage=True,
            showErrorMessage=False,
            promptTitle=(col.label or col.code)[:32],
            prompt=col.prompt[:255],
        )
        letter = get_column_letter(col_idx)
        dv.add(f"{letter}3:{letter}{DATA_ROWS_END}")
        ws.add_data_validation(dv)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
